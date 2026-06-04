import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Define target paths
OUTPUT_DIR = r"d:\DEV\conf-platform\excel_templates"
os.makedirs(OUTPUT_DIR, exist_ok=True)

agenda_path = os.path.join(OUTPUT_DIR, "conference_agenda_template.xlsx")
registration_path = os.path.join(OUTPUT_DIR, "registration_roster_template.xlsx")

def create_agenda_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"  # Sheet name must be Schedule or the first sheet
    
    # Define columns
    headers = [
        "Session Code", "Session Name", "Session Type", "Room Name",
        "Start Datetime", "End Datetime", "Speaker First Name", "Speaker Last Name",
        "Speaker Email", "Speaker Phone", "Speaker Affiliation", "Speaker Country",
        "Speaker Designation", "Presentation Title", "Talk Start", "Talk End",
        "Talk Order", "Talk Duration", "Moderator Name"
    ]
    
    ws.append(headers)
    
    # Sample Data
    data = [
        [
            "KEY-01", "Opening Keynote: The Future of Conference Tech", "keynote", "Main Auditorium",
            "2026-10-15 09:00", "2026-10-15 10:00", "Sarah", "Chen",
            "sarah.chen@techfuture.org", "+1-555-0199", "TechFuture Foundation", "United States",
            "Chief Technology Officer", "Antigravity AI: A New Paradigm in Coding", "09:05", "09:50",
            1, 45, "Dr. Alan Turing"
        ],
        [
            "TECH-02", "Web Engineering & Scale", "regular", "Hall B",
            "2026-10-15 10:30", "2026-10-15 12:00", "Michael", "O'Connor",
            "moconnor@scaleops.io", "+353-1-496-0123", "ScaleOps Ireland", "Ireland",
            "Principal Infrastructure Engineer", "Optimizing Next.js Turbopack at Enterprise Scale", "10:35", "11:15",
            1, 40, "Grace Hopper"
        ],
        [
            "TECH-02", "Web Engineering & Scale", "regular", "Hall B",
            "2026-10-15 10:30", "2026-10-15 12:00", "Yuki", "Tanaka",
            "y.tanaka@tokyoweb.jp", "+81-3-5555-0143", "Tokyo Web Lab", "Japan",
            "Senior Software Architect", "Real-time Collaboration Patterns with FastAPI", "11:20", "11:55",
            2, 35, "Grace Hopper"
        ]
    ]
    
    for row in data:
        ws.append(row)
        
    # Styling
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10)
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_left = Alignment(horizontal="left", vertical="center")
    thin_border_side = Side(border_style="thin", color="CCCCCC")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # Apply to headers
    ws.row_dimensions[1].height = 28
    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = alignment_center
        cell.border = thin_border
        
    # Apply to data rows
    for r_idx in range(2, len(data) + 2):
        ws.row_dimensions[r_idx].height = 22
        for c_idx, cell in enumerate(ws[r_idx], 1):
            cell.font = font_data
            cell.border = thin_border
            # Left align text, center codes and datetimes/times
            if c_idx in [1, 3, 5, 6, 15, 16, 17, 18]:
                cell.alignment = alignment_center
            else:
                cell.alignment = alignment_left
                
    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    wb.save(agenda_path)
    print(f"Created Speaker Agenda template at: {agenda_path}")

def create_registration_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    
    # Define columns
    headers = [
        "First Name *", "Last Name *", "Email Address *", "Phone Number",
        "Company/Affiliation", "Job Title/Designation", "Country",
        "Registration Category *", "Paid Status", "Source"
    ]
    
    ws.append(headers)
    
    # Sample Data
    data = [
        [
            "Asha", "Mehta", "asha@example.com", "+91 9876543210",
            "Asha Industries", "Managing Director", "India",
            "Delegate", "Unpaid", "excel_import"
        ],
        [
            "John", "Doe", "johndoe@example.com", "+1-555-0148",
            "Globex Corp", "Director of Product", "Canada",
            "VIP", "Paid", "excel_import"
        ]
    ]
    
    for row in data:
        ws.append(row)
        
    # Styling
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10)
    fill_header = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")  # Emerald Green
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_left = Alignment(horizontal="left", vertical="center")
    thin_border_side = Side(border_style="thin", color="CCCCCC")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # Apply to headers
    ws.row_dimensions[1].height = 28
    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = alignment_center
        cell.border = thin_border
        
    # Apply to data rows
    for r_idx in range(2, len(data) + 2):
        ws.row_dimensions[r_idx].height = 22
        for c_idx, cell in enumerate(ws[r_idx], 1):
            cell.font = font_data
            cell.border = thin_border
            if c_idx in [3, 4, 7, 8, 9, 10]:
                cell.alignment = alignment_center
            else:
                cell.alignment = alignment_left
                
    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    wb.save(registration_path)
    print(f"Created Registration Roster template at: {registration_path}")

if __name__ == "__main__":
    create_agenda_template()
    create_registration_template()
