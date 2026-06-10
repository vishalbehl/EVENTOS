"""
inspect_db.py  —  PostgreSQL schema → Word document
────────────────────────────────────────────────────
Usage:
    pip install psycopg2-binary python-docx
    python inspect_db.py

    DB_URL="postgresql://user:pass@host:5432/db" python inspect_db.py

Output:  db_schema_<dbname>.docx
"""

import os, sys, json, datetime, re
import psycopg2
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK

# ── Config ─────────────────────────────────────────────────────────────────────

DB_URL = os.getenv("DB_URL", "postgresql://postgres:847425@localhost:5432/eventos_db")

CLR_HDR_BG = "1F3864"
CLR_HDR_FG = "FFFFFF"
CLR_ALT    = "EBF3FB"
CLR_ACCENT = "2E75B6"
CLR_BORDER = "BFBFBF"
CLR_NAVY   = "1F3864"
CLR_GREY   = "595959"

# ── XML helpers ────────────────────────────────────────────────────────────────

def _cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def _cell_borders(cell, color=CLR_BORDER):
    tcPr = cell._tc.get_or_add_tcPr()
    b    = OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        b.append(el)
    tcPr.append(b)

def _cell_padding(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    mar  = OxmlElement("w:tcMar")
    for side, val in (("top","60"),("bottom","60"),("left","100"),("right","100")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"),    val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)

def _spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr()
    spg = OxmlElement("w:spacing")
    spg.set(qn("w:before"), str(before))
    spg.set(qn("w:after"),  str(after))
    pPr.append(spg)

def _write_cell(cell, text, bold=False, italic=False, size_pt=9,
                color=None, align=None, bg=None, border_color=CLR_BORDER):
    if bg:
        _cell_bg(cell, bg)
    _cell_borders(cell, border_color)
    _cell_padding(cell)
    p = cell.paragraphs[0]
    p.clear()
    if align:
        p.alignment = align
    _spacing(p)
    run = p.add_run(str(text))
    run.bold      = bold
    run.italic    = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def _add_para(doc, text="", bold=False, italic=False, size_pt=10,
              color=None, before=40, after=40, align=None):
    p = doc.add_paragraph()
    _spacing(p, before=before, after=after)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(str(text))
        run.bold      = bold
        run.italic    = italic
        run.font.name = "Calibri"
        run.font.size = Pt(size_pt)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p

def _page_break(doc):
    """Zero-height paragraph containing a page break — no blank page side-effects."""
    p   = doc.add_paragraph()
    _spacing(p, before=0, after=0)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

# ── DB helpers ─────────────────────────────────────────────────────────────────

def db_parent_tables(cur):
    """
    Return only the top-level (non-partition) tables grouped by schema.
    """
    # All base tables across all user-defined schemas
    cur.execute("""
        SELECT table_schema, table_name
        FROM   information_schema.tables
        WHERE  table_schema NOT IN ('pg_catalog', 'information_schema')
          AND  table_schema NOT LIKE 'pg_toast%'
          AND  table_schema NOT LIKE 'pg_temp%'
          AND  table_type   = 'BASE TABLE'
        ORDER  BY table_schema, table_name
    """)
    all_tables = {r for r in cur.fetchall()}

    # Tables that are children in pg_inherits
    cur.execute("""
        SELECT n.nspname, c.relname
        FROM   pg_inherits i
        JOIN   pg_class    c ON c.oid = i.inhrelid
        JOIN   pg_namespace n ON n.oid = c.relnamespace
    """)
    child_tables = {r for r in cur.fetchall()}

    # Regex fallback for any missed partition naming patterns
    partition_re = re.compile(
        r'(_y\d{4}m\d{2}|_default|_p\d+)$', re.IGNORECASE
    )

    result = []
    for schema, table in sorted(all_tables):
        if (schema, table) in child_tables:
            continue
        if partition_re.search(table):
            continue
        result.append((schema, table))

    return result

def db_columns(cur, schema, table):
    cur.execute("""
        SELECT c.column_name, c.data_type,
               c.character_maximum_length, c.is_nullable,
               c.column_default,
               pgd.description
        FROM   information_schema.columns c
        LEFT   JOIN pg_class       pgc ON pgc.relname  = c.table_name
                                      AND pgc.relnamespace = (
                                          SELECT oid FROM pg_namespace WHERE nspname = c.table_schema
                                      )
        LEFT   JOIN pg_description pgd ON pgd.objoid   = pgc.oid
                                      AND pgd.objsubid = c.ordinal_position
        WHERE  c.table_schema = %s AND c.table_name = %s
        ORDER  BY c.ordinal_position
    """, (schema, table))
    return cur.fetchall()

def db_constraints(cur, schema, table):
    cur.execute("""
        SELECT kcu.column_name, tc.constraint_type
        FROM   information_schema.table_constraints tc
        JOIN   information_schema.key_column_usage  kcu
               ON  tc.constraint_name = kcu.constraint_name
               AND tc.table_schema    = kcu.table_schema
        WHERE  tc.table_schema = %s AND tc.table_name = %s
    """, (schema, table))
    out = {}
    for col, ctype in cur.fetchall():
        out.setdefault(col, []).append(ctype)
    return out

def db_sample(cur, schema, table):
    try:
        cur.execute(f'SELECT * FROM "{schema}"."{table}" LIMIT 1')
        cols = [d[0] for d in cur.description]
        row  = cur.fetchone()
        return (cols, list(row)) if row else ([], [])
    except Exception:
        return [], []

def db_count(cur, schema, table):
    try:
        cur.execute(f'SELECT COUNT(*) FROM "{schema}"."{table}"')
        return cur.fetchone()[0]
    except Exception:
        return "?"

def fmt(val):
    if val is None:
        return "NULL"
    if isinstance(val, dict):
        return json.dumps(val, default=str, ensure_ascii=False)[:160]
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.isoformat()
    return str(val)[:160]

# ── Document builder ───────────────────────────────────────────────────────────

def build(tables_data, db_name):
    doc = Document()

    # Kill default paragraph spacing on Normal
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    nPr = normal._element.get_or_add_pPr()
    spg = OxmlElement("w:spacing")
    spg.set(qn("w:before"), "0")
    spg.set(qn("w:after"),  "0")
    nPr.append(spg)

    # A4 portrait, tight margins
    sec = doc.sections[0]
    sec.page_width   = Inches(8.3)
    sec.page_height  = Inches(11.7)
    sec.left_margin  = sec.right_margin  = Inches(0.6)
    sec.top_margin   = sec.bottom_margin = Inches(0.6)
    content_w = 8.3 - 2 * 0.6   # 7.1 inches

    # Column widths for schema table  [#, Name, Type, Null, Constraint, Default]
    C_W = [0.3, 1.4, 1.2, 0.5, 1.0,
       content_w - 0.3 - 1.4 - 1.2 - 0.5 - 1.0]    
    HDR = ["#", "Column Name", "Data Type", "Nullable", "Constraint", "Default"]

    # Unique schemas list
    schemas_list = sorted(list({td["schema"] for td in tables_data}))

    # ── Cover ──
    _add_para(doc, "Database Schema Reference",
              bold=True, size_pt=22, color=CLR_NAVY,
              before=100, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc,
              f"Database: {db_name}   ·   "
              f"Generated: {datetime.datetime.now().strftime('%d %b %Y %H:%M')}   ·   "
              f"Tables: {len(tables_data)}\n"
              f"Schemas: {', '.join(schemas_list)}",
              italic=True, size_pt=10, color=CLR_GREY,
              before=0, after=60, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ── TOC — two-column ──
    _add_para(doc, "Table of Contents",
              bold=True, size_pt=13, color=CLR_NAVY, before=40, after=20)

    half        = (len(tables_data) + 1) // 2
    left_items  = tables_data[:half]
    right_items = tables_data[half:]
    toc_tbl     = doc.add_table(rows=half, cols=2)
    toc_w       = int(Inches(content_w / 2))

    for ri in range(half):
        for ci, items in enumerate([left_items, right_items]):
            cell = toc_tbl.rows[ri].cells[ci]
            cell.width = toc_w
            _cell_padding(cell)
            p = cell.paragraphs[0]
            _spacing(p)
            if ri < len(items):
                td  = items[ri]
                idx = (ri if ci == 0 else ri + half) + 1
                r1  = p.add_run(f"{idx}. {td['schema']}.{td['table']}")
                r1.font.name = "Calibri"
                r1.font.size = Pt(9)
                r1.bold = True
                r2  = p.add_run(f"  ({td['row_count']} rows, {len(td['columns'])} cols)")
                r2.font.name = "Calibri"
                r2.font.size = Pt(8)
                r2.italic = True
                r2.font.color.rgb = RGBColor.from_string(CLR_GREY)

    # ── One section per table ──
    for idx, td in enumerate(tables_data, 1):

        # Page break before every table (first table also gets one,
        # which pushes it past the TOC — clean and predictable)
        _page_break(doc)

        schema = td["schema"]
        tname  = td["table"]
        cols   = td["columns"]
        constr = td["constraints"]
        s_cols = td["sample_cols"]
        s_vals = td["sample_vals"]
        rcount = td["row_count"]

        # Heading
        _add_para(doc, f"{idx}. {schema}.{tname}",
                  bold=True, size_pt=13, color=CLR_NAVY, before=0, after=10)
        _add_para(doc, f"Schema: {schema}   ·   Rows: {rcount}   ·   Columns: {len(cols)}",
                  italic=True, size_pt=8, color=CLR_GREY, before=0, after=20)

        # Schema table
        tbl = doc.add_table(rows=1 + len(cols), cols=len(HDR))
        tbl.style = "Table Grid"

        for ci, (h, w) in enumerate(zip(HDR, C_W)):
            cell = tbl.rows[0].cells[ci]
            cell.width = Inches(w)
            _write_cell(cell, h, bold=True, size_pt=8,
                        color=CLR_HDR_FG, bg=CLR_HDR_BG,
                        border_color=CLR_HDR_BG,
                        align=WD_ALIGN_PARAGRAPH.CENTER)

        for ri, col in enumerate(cols):
            col_name, dtype, max_len, nullable, default, _ = col
            bg   = CLR_ALT if ri % 2 == 0 else "FFFFFF"
            tstr = dtype + (f"({max_len})" if max_len else "")
            dstr = (str(default) if default else "")[:55]
            cstr = ", ".join(constr.get(col_name, []))
            vals = [str(ri+1), col_name, tstr,
                    "YES" if nullable == "YES" else "NO",
                    cstr, dstr]
            for ci, (v, w) in enumerate(zip(vals, C_W)):
                cell = tbl.rows[ri+1].cells[ci]
                cell.width = Inches(w)
                _write_cell(cell, v, bold=(ci == 1), size_pt=8, bg=bg)

        # Sample row
        _add_para(doc, "Sample Row  (1 real record from DB)",
                  bold=True, size_pt=9, color=CLR_ACCENT, before=25, after=8)

        if s_cols:
            s_tbl = doc.add_table(rows=1 + len(s_cols), cols=2)
            s_tbl.style = "Table Grid"

            for ci, h in enumerate(["Column", "Value"]):
                w    = 2.0 if ci == 0 else content_w - 2.0
                cell = s_tbl.rows[0].cells[ci]
                cell.width = Inches(w)
                _write_cell(cell, h, bold=True, size_pt=8,
                            color=CLR_HDR_FG, bg=CLR_ACCENT,
                            border_color=CLR_ACCENT,
                            align=WD_ALIGN_PARAGRAPH.CENTER)

            for ri, (sc, sv) in enumerate(zip(s_cols, s_vals)):
                bg   = CLR_ALT if ri % 2 == 0 else "FFFFFF"
                row  = s_tbl.rows[ri+1]
                cell0 = row.cells[0]
                cell0.width = Inches(2.0)
                _write_cell(cell0, sc, bold=True, size_pt=8, bg=bg)
                cell1 = row.cells[1]
                cell1.width = Inches(content_w - 2.0)
                _write_cell(cell1, fmt(sv), size_pt=8, bg=bg)
        else:
            _add_para(doc, "No rows in this table.",
                      italic=True, size_pt=8, color="999999", before=0, after=0)

    return doc

# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    print(f"Connecting: {DB_URL}")
    try:
        conn = psycopg2.connect(DB_URL)
        conn.set_session(readonly=True)
    except Exception as e:
        print(f"Connection failed: {e}")
        sys.exit(1)

    cur     = conn.cursor()
    db_name = DB_URL.rstrip("/").split("/")[-1]

    print("Reading parent tables across all schemas (partitions excluded)…")
    tables = db_parent_tables(cur)
    print(f"Found {len(tables)} tables (after filtering partitions).\n")

    data = []
    for s, t in tables:
        print(f"  {s}.{t}")
        s_cols, s_vals = db_sample(cur, s, t)
        data.append({
            "schema":      s,
            "table":       t,
            "columns":     db_columns(cur, s, t),
            "constraints": db_constraints(cur, s, t),
            "sample_cols": s_cols,
            "sample_vals": s_vals,
            "row_count":   db_count(cur, s, t),
        })
    conn.close()

    print("\nBuilding document…")
    doc = build(data, db_name)
    out = f"db_schema_{db_name}.docx"
    doc.save(out)
    print(f"\nSaved: {out}  ({len(data)} tables across user schemas, partitions excluded)")

if __name__ == "__main__":
    main()