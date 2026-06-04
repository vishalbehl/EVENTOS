
import openpyxl
from openpyxl.styles import Font, PatternFill

def generate_sample():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"

    headers = [
        "session_code", "session_name", "session_type", "room_name", 
        "start_datetime", "end_datetime", 
        "speaker_first_name", "speaker_last_name", "speaker_email",
        "speaker_phone", "speaker_affiliation", "speaker_country",
        "presentation_title", "talk_start_time", "talk_end_time",
        "talk_order", "talk_duration_min", "moderator_name"
    ]

    # Style headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill

    # Sample Data
    data = [
        ["S-101", "Opening Plenary", "regular", "Hall A", "2026-09-01 09:00", "2026-09-01 10:30", "John", "Doe", "john@example.com", "+123456789", "Tech University", "USA", "The Future of AI", "09:00", "09:20", 0, 20, "Jane Smith"],
        ["S-101", "Opening Plenary", "regular", "Hall A", "2026-09-01 09:00", "2026-09-01 10:30", "Alice", "Brown", "alice@example.com", None, "Global Systems", "UK", "Quantum Computing Basics", "09:20", "09:40", 1, 20, "Jane Smith"],
        ["S-202", "Workshop: Next.js", "workshop", "Room 302", "2026-09-01 11:00", "2026-09-01 13:00", "Bob", "Wilson", "bob@example.com", None, "Vercel", "USA", "Deep Dive into App Router", "11:00", "12:30", 0, 90, "Mark Lee"],
    ]

    for row_num, row_data in enumerate(data, 2):
        for col_num, value in enumerate(row_data, 1):
            ws.cell(row=row_num, column=col_num).value = value

    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save("apps/cloud/command-center/public/samples/sample_agenda.xlsx")
    print("Sample agenda generated successfully.")

if __name__ == "__main__":
    generate_sample()
