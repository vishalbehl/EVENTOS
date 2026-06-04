from openpyxl import Workbook
from faker import Faker
import random

fake = Faker()

wb = Workbook()
ws = wb.active
ws.title = "Participants"

headers = [
    "First Name *",
    "Last Name *",
    "Email Address *",
    "Phone Number",
    "Company/Affiliation",
    "Job Title/Designation",
    "Country",
    "Registration Category *",
    "Paid Status"
]

for c,h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)

roles = [
    ("Delegate","DEL"),
    ("Student Delegate","STU"),
    ("Organizer","ORG"),
    ("Speaker","SPK"),
    ("Sponsor Representative","SPO"),
    ("Exhibitor","EXH"),
    ("Media","MED"),
    ("Technical Staff","TEC"),
    ("VIP Guest","VIP"),
]

countries = {
    "India":"+91","United States":"+1","United Kingdom":"+44","Germany":"+49",
    "France":"+33","Canada":"+1","Australia":"+61","Singapore":"+65",
    "Japan":"+81","UAE":"+971","Netherlands":"+31","Sweden":"+46"
}

used=set()

for r in range(2,502):
    first = fake.first_name()
    last = fake.last_name()

    email = f"{first.lower()}{last.lower()}{random.randint(10,99)}@example.com"
    while email in used:
        email = f"{first.lower()}{last.lower()}{random.randint(10,99)}@example.com"
    used.add(email)

    country = random.choice(list(countries.keys()))
    phone = countries[country] + ''.join(random.choices('0123456789', k=9))

    role_name, role_code = random.choice(roles)

    ws.cell(r,1,first)
    ws.cell(r,2,last)
    ws.cell(r,3,email)
    ws.cell(r,4,phone)
    ws.cell(r,5,fake.company())
    ws.cell(r,6,random.choice([
        "Manager","Director","Engineer","Consultant","Researcher",
        "Coordinator","Analyst","Executive","Lead","Professor"
    ]))
    ws.cell(r,7,country)
    ws.cell(r,8,f"{role_name}")
    ws.cell(r,9,"Unpaid")

path="./registration_roster_500_corrected.xlsx"
wb.save(path)

print(path)
