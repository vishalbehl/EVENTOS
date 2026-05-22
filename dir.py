import os
from pathlib import Path
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate, Preformatted
from reportlab.lib import styles

# ======================================
# Configuration
# ======================================

PROJECT_ROOT = Path.cwd()

EXCLUDED_DIRS = {
    "node_modules",
    "__pycache__",
    ".git",
    ".next",
    "dist",
    "build",
    ".venv",
    "venv",
    "env",
    "Scripts",
    "Lib",
    "site-packages",
    ".idea",
    ".vscode",
    ".pytest_cache",
    ".mypy_cache",
    ".cache",
    "coverage",
    ".turbo",
    ".parcel-cache",
    ".DS_Store"
}

EXCLUDED_EXTENSIONS = {
    ".pyc",
    ".pyo",
    ".log",
    ".tmp",
    ".cache"
}

EXCLUDED_FILES = {
    "package-lock.json",
    "yarn.lock",
    "pnpm-lock.yaml",
    ".env",
    ".env.local"
}

# ======================================
# Build Directory Tree
# ======================================

def generate_tree(directory, prefix=""):
    tree = []

    items = sorted(
        directory.iterdir(),
        key=lambda x: (x.is_file(), x.name.lower())
    )

    filtered = []

    for item in items:
        if item.name in EXCLUDED_DIRS:
            continue

        if item.name in EXCLUDED_FILES:
            continue

        if item.suffix.lower() in EXCLUDED_EXTENSIONS:
            continue

        filtered.append(item)

    for i, item in enumerate(filtered):
        connector = "└── " if i == len(filtered)-1 else "├── "

        tree.append(prefix + connector + item.name)

        if item.is_dir():
            extension = "    " if i == len(filtered)-1 else "│   "
            tree.extend(
                generate_tree(
                    item,
                    prefix + extension
                )
            )

    return tree


# ======================================
# Generate Tree Text
# ======================================

project_name = PROJECT_ROOT.name

tree_output = [project_name]
tree_output.extend(generate_tree(PROJECT_ROOT))

tree_text = "\n".join(tree_output)

print("\nGenerated directory structure:\n")
print(tree_text)

# ======================================
# Save DOCX
# ======================================

doc = Document()

doc.add_heading(
    f"Project Directory Structure - {project_name}",
    level=1
)

p = doc.add_paragraph()
run = p.add_run(tree_text)

run.font.name = "Courier New"
run.font.size = Pt(8)

doc.save("project_structure.docx")

print("\nDOCX created:")
print("project_structure.docx")


# ======================================
# Save PDF
# ======================================

pdf = SimpleDocTemplate(
    "project_structure.pdf"
)

style = styles.getSampleStyleSheet()

mono_style = style["Code"]
mono_style.fontName = "Courier"
mono_style.fontSize = 7
mono_style.leading = 8

content = [
    Preformatted(
        tree_text,
        mono_style
    )
]

pdf.build(content)

print("PDF created:")
print("project_structure.pdf")


print("\nDone.")