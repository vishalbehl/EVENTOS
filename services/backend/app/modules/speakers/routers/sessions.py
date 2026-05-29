# backend/app/routers/sessions.py
from __future__ import annotations

import uuid
from typing import List, Optional, Annotated

from fastapi import APIRouter, Depends, HTTPException, Query, status
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_db, get_current_event, CurrentEvent, OrganizerOrAbove, get_current_user
from app.modules.speakers.models.session import Session
from app.modules.speakers.models.session_speaker import SessionSpeaker
from app.modules.auth.models.user import User
from app.modules.speakers.schemas.session import (
    SessionCreate, SessionUpdate, SessionResponse, SessionSummary,
    SessionSpeakerCreate, ReorderSpeakersRequest, SessionSpeakerUpdate,
)
from app.schemas.common import MessageResponse

router = APIRouter(prefix="/events/{event_id}/sessions", tags=["sessions"])


@router.get("/export")
async def export_sessions_docx(
    event: CurrentEvent,
    db: AsyncSession = Depends(get_db),
):
    """
    Export all event sessions as a beautifully formatted Word (.docx) agenda.
    """
    import io
    from datetime import datetime
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from fastapi.responses import StreamingResponse
    
    # 1. Fetch all sessions for this event sorted by start time
    stmt = (
        select(Session)
        .options(
            selectinload(Session.room),
            selectinload(Session.session_speakers).selectinload(SessionSpeaker.speaker)
        )
        .where(Session.event_id == event.id)
        .order_by(Session.start_time)
    )
    res = await db.execute(stmt)
    sessions = res.scalars().all()

    # 1.1 Localize timezone setup
    from zoneinfo import ZoneInfo
    from app.services.timezone_service import get_cached_timezone
    from datetime import timezone
    tz_name = event.timezone or get_cached_timezone()
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = ZoneInfo("Asia/Kolkata")

    def localize_dt(dt):
        if not dt:
            return None
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz)

    # Group sessions by date
    from collections import defaultdict
    sessions_by_date = defaultdict(list)
    for s in sessions:
        local_start = localize_dt(s.start_time)
        date_str = local_start.strftime("%A, %d %B %Y") if local_start else "TBD"
        sessions_by_date[date_str].append(s)

    # 2. Build Document
    doc = Document()

    # Style definitions
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)

    # Title block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(f"OFFICIAL AGENDA & SCHEDULE")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(49, 46, 129)  # Deep Indigo (#312e81)

    event_p = doc.add_paragraph()
    event_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    event_run = event_p.add_run(event.name)
    event_run.font.size = Pt(15)
    event_run.font.bold = True
    event_run.font.color.rgb = RGBColor(79, 70, 229)  # Vibrant Indigo (#4f46e5)

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    loc = event.location or "Venue Location"
    info_run = info_p.add_run(f"Location: {loc}\nDate generated: {datetime.now().strftime('%d %B %Y')}")
    info_run.font.size = Pt(10)
    info_run.font.italic = True
    info_run.font.color.rgb = RGBColor(100, 116, 139)  # Slate gray

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Daily agendas
    for date_str, daily_sessions in sorted(sessions_by_date.items(), key=lambda x: x[1][0].start_time):
        doc.add_heading(date_str, level=1)
        h1 = doc.paragraphs[-1]
        h1.runs[0].font.name = 'Arial'
        h1.runs[0].font.size = Pt(15)
        h1.runs[0].font.bold = True
        h1.runs[0].font.color.rgb = RGBColor(49, 46, 129)

        # Add a table for this day's schedule
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        table.allow_autofit = False

        # Width definitions
        widths = [Inches(1.5), Inches(1.5), Inches(4.5)]
        
        # Header row styling
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Time Block'
        hdr_cells[1].text = 'Session & Room'
        hdr_cells[2].text = 'Agenda Details / Presentation Topics'

        for i, cell in enumerate(hdr_cells):
            cell.width = widths[i]
            # Set background color of header
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="312e81"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(10.5)

        for session in sorted(daily_sessions, key=lambda x: x.start_time):
            row_cells = table.add_row().cells
            for i, cell in enumerate(row_cells):
                cell.width = widths[i]

            local_start = localize_dt(session.start_time)
            local_end = localize_dt(session.end_time)
            if local_start and local_end:
                time_range = f"{local_start.strftime('%I:%M %p')} - {local_end.strftime('%I:%M %p')}"
            else:
                time_range = "TBD"
            row_cells[0].text = time_range
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            row_cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)

            room_name = session.room.name if session.room else "Unassigned Room"
            row_cells[1].text = f"Code: {session.session_code}\nRoom: {room_name}"
            for r in row_cells[1].paragraphs[0].runs:
                r.font.size = Pt(9.5)
            row_cells[1].paragraphs[0].runs[0].font.bold = True

            # Agenda details: Session title, description, speakers
            p_desc = row_cells[2].paragraphs[0]
            p_title_run = p_desc.add_run(session.name)
            p_title_run.font.bold = True
            p_title_run.font.size = Pt(11)
            p_title_run.font.color.rgb = RGBColor(79, 70, 229)

            if session.description:
                p_desc.add_run(f"\n{session.description}").font.size = Pt(9)
                p_desc.runs[-1].font.color.rgb = RGBColor(100, 116, 139)

            if session.session_speakers:
                p_desc.add_run("\n\nPresentations:").font.bold = True
                p_desc.runs[-1].font.size = Pt(9.5)
                for ss in sorted(session.session_speakers, key=lambda x: x.talk_order):
                    sp_name = ss.speaker.full_name if ss.speaker else "Unknown Speaker"
                    title = ss.presentation_title or "No Title Provided"
                    duration = f"{ss.talk_duration_minutes} min" if ss.talk_duration_minutes else ""
                    time_info = ""
                    if ss.start_time and ss.end_time:
                        local_ss_start = localize_dt(ss.start_time)
                        local_ss_end = localize_dt(ss.end_time)
                        if local_ss_start and local_ss_end:
                            time_info = f" ({local_ss_start.strftime('%I:%M %p')} - {local_ss_end.strftime('%I:%M %p')})"
                    
                    bullets_p = row_cells[2].add_paragraph(style='List Bullet')
                    bullets_p.paragraph_format.space_before = Pt(2)
                    bullets_p.paragraph_format.space_after = Pt(2)
                    
                    r_sp = bullets_p.add_run(f"{sp_name}")
                    r_sp.font.bold = True
                    r_sp.font.size = Pt(9.5)
                    
                    r_title = bullets_p.add_run(f" — \"{title}\"")
                    r_title.font.italic = True
                    r_title.font.size = Pt(9.5)
                    
                    if duration or time_info:
                        r_time = bullets_p.add_run(f" [{duration}{time_info}]")
                        r_time.font.size = Pt(8.5)
                        r_time.font.color.rgb = RGBColor(100, 116, 139)

            # Cell Margins padding
            for cell in row_cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = OxmlElement('w:tcMar')
                for margin_type in ['top', 'bottom', 'left', 'right']:
                    node = OxmlElement(f'w:{margin_type}')
                    node.set(qn('w:w'), '120')
                    node.set(qn('w:type'), 'dxa')
                    tcMar.append(node)
                tcPr.append(tcMar)

        doc.add_paragraph().paragraph_format.space_after = Pt(18)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    filename = f"agenda-{event.short_code.lower()}-{datetime.now().strftime('%Y%m%d')}.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("", response_model=List[SessionSummary])
async def list_sessions(
    event: CurrentEvent,
    current_user: Annotated[User, Depends(get_current_user)],
    db: AsyncSession = Depends(get_db),
    room_id: Optional[uuid.UUID] = Query(None),
    status_filter: Optional[str] = Query(None, alias="status"),
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=1000),
) -> List[SessionSummary]:
    q = select(Session).options(
        selectinload(Session.event), 
        selectinload(Session.room),
        selectinload(Session.session_speakers).selectinload(SessionSpeaker.speaker),
        selectinload(Session.posters)
    ).where(Session.event_id == event.id)

    # Restricted roles (NOT super_admin or admin) must have specific assignments
    if current_user.role not in ["super_admin", "admin"]:
        from app.modules.rbac.models.rbac import UserAccessNode
        from sqlalchemy import or_, exists

        # Check Event Level Assignment
        event_assigned = exists().where(
            UserAccessNode.user_id == current_user.id,
            UserAccessNode.node_id == event.id,
            UserAccessNode.node_type == 'EVENT'
        )

        # Check Room Level Assignment (If session is in an assigned room)
        room_assignments = select(UserAccessNode.node_id).where(
            UserAccessNode.user_id == current_user.id,
            UserAccessNode.node_type == 'ROOM'
        )

        # Check Session Level Assignment
        session_assignments = select(UserAccessNode.node_id).where(
            UserAccessNode.user_id == current_user.id,
            UserAccessNode.node_type == 'SESSION'
        )

        q = q.where(
            or_(
                event_assigned,
                Session.room_id.in_(room_assignments),
                Session.id.in_(session_assignments)
            )
        )

    if room_id:
        q = q.where(Session.room_id == room_id)
    if status_filter:
        q = q.where(Session.status == status_filter)
    q = q.order_by(Session.start_time).offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(q)
    return [SessionSummary.model_validate(s) for s in result.scalars().all()]


@router.post("", response_model=SessionResponse, status_code=status.HTTP_201_CREATED)
async def create_session(
    payload: SessionCreate,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> SessionResponse:
    # Enforce assignments for restricted roles
    if current_user.role not in ["super_admin", "admin", "organiser"]:
        from app.modules.rbac.models.rbac import UserAccessNode
        from sqlalchemy import or_, and_

        # Check if assigned to the event, or the specific room in the payload
        assignment_check = await db.execute(
            select(UserAccessNode).where(
                UserAccessNode.user_id == current_user.id,
                or_(
                    and_(UserAccessNode.node_id == event.id, UserAccessNode.node_type == 'EVENT'),
                    and_(UserAccessNode.node_id == payload.room_id, UserAccessNode.node_type == 'ROOM')
                )
            )
        )
        if not assignment_check.scalar_one_or_none():
            # Check legacy assignments as fallback
            from app.modules.rbac.models.user_assignment import UserEventAssignment
            legacy_check = await db.execute(
                select(UserEventAssignment).where(
                    UserEventAssignment.user_id == current_user.id,
                    UserEventAssignment.event_id == event.id
                )
            )
            if not legacy_check.scalar_one_or_none():
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN, 
                    detail="You do not have permission to create sessions in this room."
                )

    """Register a new session."""
    dup = await db.execute(
        select(Session).where(
            Session.event_id == event.id,
            Session.session_code == payload.session_code,
        )
    )
    if dup.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_409_CONFLICT,
                            detail=f"Session code '{payload.session_code}' already used.")

    data = payload.model_dump()
    speakers_data = data.pop("speakers", None) or []
    session = Session(event_id=event.id, **data)
    db.add(session)
    await db.flush()

    for idx, sp_data in enumerate(speakers_data):
        ss = SessionSpeaker(
            session_id=session.id,
            speaker_id=sp_data["speaker_id"],
            presentation_title=sp_data.get("presentation_title"),
            talk_order=sp_data.get("talk_order", idx),
            talk_duration_minutes=sp_data.get("talk_duration_minutes"),
            start_time=sp_data.get("start_time"),
            end_time=sp_data.get("end_time"),
        )
        db.add(ss)

    await db.commit()
    return SessionResponse.model_validate(
        await _get_session_or_404(db, session.id, event.id)
    )


@router.get("/{session_id}", response_model=SessionResponse)
async def get_session(
    session_id: uuid.UUID,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> SessionResponse:
    return SessionResponse.model_validate(
        await _get_session_or_404(db, session_id, event.id, user=current_user)
    )


@router.patch("/{session_id}", response_model=SessionResponse)
async def update_session(
    session_id: uuid.UUID,
    payload: SessionUpdate,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> SessionResponse:
    session = await _get_session_or_404(db, session_id, event.id, user=current_user)
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(session, field, value)
    await db.commit()
    return SessionResponse.model_validate(
        await _get_session_or_404(db, session.id, event.id, user=current_user)
    )


@router.delete("/{session_id}", response_model=MessageResponse)
async def delete_session(
    session_id: uuid.UUID,
    event: CurrentEvent,
    user: OrganizerOrAbove,
    db: AsyncSession = Depends(get_db),
) -> MessageResponse:
    session = await _get_session_or_404(db, session_id, event.id, user=user)
    if session.status == "in_progress":
        raise HTTPException(status_code=status.HTTP_409_CONFLICT,
                            detail="Cannot delete a session that is in progress.")
    await db.delete(session)
    await db.commit()
    return MessageResponse(message="Session deleted.")


@router.post("/{session_id}/speakers", response_model=MessageResponse)
async def add_speaker_to_session(
    session_id: uuid.UUID,
    payload: SessionSpeakerCreate,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MessageResponse:
    await _get_session_or_404(db, session_id, event.id, user=current_user)
    dup = await db.execute(
        select(SessionSpeaker).where(
            SessionSpeaker.session_id == session_id,
            SessionSpeaker.speaker_id == payload.speaker_id,
        )
    )
    if dup.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_409_CONFLICT,
                            detail="Speaker already linked to this session.")
    ss = SessionSpeaker(
        session_id=session_id,
        speaker_id=payload.speaker_id,
        presentation_title=payload.presentation_title,
        talk_order=payload.talk_order or 0,
        talk_duration_minutes=payload.talk_duration_minutes,
        start_time=payload.start_time,
        end_time=payload.end_time,
    )
    db.add(ss)
    await db.commit()
    return MessageResponse(message="Speaker added to session.")


@router.delete("/{session_id}/speakers/{session_speaker_id}", response_model=MessageResponse)
async def remove_speaker_from_session(
    session_id: uuid.UUID,
    session_speaker_id: uuid.UUID,
    event: CurrentEvent,
    user: OrganizerOrAbove,
    db: AsyncSession = Depends(get_db),
) -> MessageResponse:
    await _get_session_or_404(db, session_id, event.id, user=user)
    result = await db.execute(
        select(SessionSpeaker).where(
            SessionSpeaker.id == session_speaker_id,
            SessionSpeaker.session_id == session_id,
        )
    )
    ss = result.scalar_one_or_none()
    if ss is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session speaker not found.")
    await db.delete(ss)
    await db.commit()
    return MessageResponse(message="Speaker removed from session.")
@router.patch("/{session_id}/speakers/{session_speaker_id}", response_model=MessageResponse)
async def update_session_speaker(
    session_id: uuid.UUID,
    session_speaker_id: uuid.UUID,
    payload: SessionSpeakerUpdate,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MessageResponse:
    # Check session access
    await _get_session_or_404(db, session_id, event.id, user=current_user)
    
    result = await db.execute(
        select(SessionSpeaker).where(
            SessionSpeaker.id == session_speaker_id,
            SessionSpeaker.session_id == session_id,
        )
    )
    ss = result.scalar_one_or_none()
    if ss is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session speaker not found.")
    
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(ss, field, value)
        
    await db.commit()
    return MessageResponse(message="Speaker talk updated/moved.")


@router.post("/{session_id}/reorder-speakers", response_model=MessageResponse)
async def reorder_speakers(
    session_id: uuid.UUID,
    payload: ReorderSpeakersRequest,
    event: CurrentEvent,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MessageResponse:
    session = await _get_session_or_404(db, session_id, event.id, user=current_user)
    current_time = session.start_time
    
    for order, ss_id in enumerate(payload.ordered_ids):
        result = await db.execute(
            select(SessionSpeaker).where(
                SessionSpeaker.id == ss_id,
                SessionSpeaker.session_id == session_id,
            )
        )
        ss = result.scalar_one_or_none()
        if ss:
            duration = ss.talk_duration_minutes or 20
            ss.talk_order = order
            ss.start_time = current_time
            from datetime import timedelta
            ss.end_time = current_time + timedelta(minutes=duration)
            current_time = ss.end_time
            
    await db.commit()
    return MessageResponse(message="Speaker order and times updated.")


async def _get_session_or_404(
    db: AsyncSession, 
    session_id: uuid.UUID, 
    event_id: uuid.UUID,
    user: Optional[User] = None
) -> Session:
    result = await db.execute(
        select(Session)
        .options(
            selectinload(Session.event),
            selectinload(Session.room),
            selectinload(Session.session_speakers).selectinload(SessionSpeaker.speaker),
            selectinload(Session.posters)
        )
        .where(Session.id == session_id, Session.event_id == event_id)
    )
    s = result.scalar_one_or_none()
    if s is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found.")

    # Enforce assignments for restricted roles
    if user and user.role not in ["super_admin", "admin", "organiser"]:
        from app.modules.rbac.models.rbac import UserAccessNode
        from sqlalchemy import or_, and_

        # Check if assigned to the event, or the specific session/room
        assignment_check = await db.execute(
            select(UserAccessNode).where(
                UserAccessNode.user_id == user.id,
                or_(
                    and_(UserAccessNode.node_id == event_id, UserAccessNode.node_type == 'EVENT'),
                    and_(UserAccessNode.node_id == session_id, UserAccessNode.node_type == 'SESSION'),
                    and_(UserAccessNode.node_id == s.room_id, UserAccessNode.node_type == 'ROOM')
                )
            )
        )
        if not assignment_check.scalars().first():
            # Check legacy assignments as fallback
            from app.modules.rbac.models.user_assignment import UserEventAssignment
            legacy_check = await db.execute(
                select(UserEventAssignment).where(
                    UserEventAssignment.user_id == user.id,
                    UserEventAssignment.event_id == event_id
                )
            )
            if not legacy_check.scalars().first():
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN, 
                    detail="You do not have permission to access this session."
                )

    return s
