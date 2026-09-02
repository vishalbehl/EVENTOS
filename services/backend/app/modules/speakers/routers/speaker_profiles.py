import asyncio
import io
import re
import uuid
from typing import List, Optional

from fastapi import APIRouter, Depends, Header, HTTPException, status, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.dependencies import get_db, get_optional_user
from app.modules.events.models.event import Event
from app.modules.events.models.speaker import Speaker
from app.modules.agenda.models import Session
from app.modules.agenda.models import SessionPerson as SessionSpeaker
from app.modules.speakers.schemas.speaker_profile import (
    SpeakerProfileCreate,
    SpeakerProfileUpdate,
    SpeakerProfileResponse,
)
from app.schemas.common import MessageResponse
from app.core.dependencies.feature_gate import enforce_event_operation
from app.modules.billing.services.usage_reservation_service import UsageReservationService
from app.modules.speakers.application.commands import SpeakerCommandService
from app.modules.speakers.application.queries import SpeakerQueryService

# Create the router with events prefix
router = APIRouter(prefix="/events/{event_id}/speakers/{speaker_id}/profile", tags=["speaker_profiles"])


async def check_profile_access(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: Optional[object] = Depends(get_optional_user),
    token: Optional[str] = Query(None),
) -> str:
    """
    Checks if the caller has access to the speaker's profile.
    Returns: 'organiser' if authorized by JWT, 'speaker' if authorized by portal token.
    Raises HTTPException (401/403/404) if unauthorized.
    """
    # 1. Try token auth first
    if token:
        q = select(Speaker).where(
            Speaker.id == speaker_id,
            Speaker.event_id == event_id,
            (Speaker.upload_token == token) | (Speaker.speaker_code == token.upper())
        ).options(selectinload(Speaker.event))
        res = await db.execute(q)
        speaker = res.scalar_one_or_none()
        if speaker:
            # Check if speaker portal is active
            if not speaker.event.speaker_mode_enabled:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Speaker portal is currently closed."
                )
            await enforce_event_operation(
                db,
                speaker.event.organization_id,
                event_id,
                "speakers.profiles.manage",
            )
            return "speaker"

    # 2. Try JWT auth
    if current_user:
        # Check if user has organizer-level access to the event
        event_res = await db.execute(select(Event).where(Event.id == event_id))
        event = event_res.scalar_one_or_none()
        if not event:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found.")

        # Check organization membership
        if getattr(current_user, "role") != "super_admin":
            if event.organization_id != getattr(current_user, "organization_id"):
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found.")

        # Check role and assignments for organisers and below
        if getattr(current_user, "role") in ["organiser", "session_manager", "technician", "volunteer"]:
            from app.modules.rbac.models.rbac import UserAccessNode
            from app.modules.agenda.models import Room
            from sqlalchemy import and_, or_
            assignment_check = await db.execute(
                select(UserAccessNode).where(
                    UserAccessNode.user_id == getattr(current_user, "id"),
                    or_(
                        and_(UserAccessNode.node_id == event_id, UserAccessNode.node_type == 'EVENT'),
                        and_(
                            UserAccessNode.node_type == 'ROOM',
                            UserAccessNode.node_id.in_(select(Room.id).where(Room.event_id == event_id))
                        ),
                        and_(
                            UserAccessNode.node_type == 'SESSION',
                            UserAccessNode.node_id.in_(select(Session.id).where(Session.event_id == event_id))
                        )
                    )
                )
            )
            if not assignment_check.scalars().first():
                from app.modules.rbac.models.user_assignment import UserEventAssignment
                legacy_check = await db.execute(
                    select(UserEventAssignment).where(
                        UserEventAssignment.user_id == getattr(current_user, "id"),
                        UserEventAssignment.event_id == event_id
                    )
                )
                if not legacy_check.scalars().first():
                    raise HTTPException(
                        status_code=status.HTTP_403_FORBIDDEN,
                        detail="You are not assigned to this event or any of its rooms/sessions."
                    )
        
        # Verify speaker exists in this event
        speaker_exists = await db.scalar(
            select(func.count(Speaker.id)).where(Speaker.id == speaker_id, Speaker.event_id == event_id)
        )
        if not speaker_exists:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Speaker not found.")
        await enforce_event_operation(
            db,
            event.organization_id,
            event_id,
            "speakers.profiles.manage",
            user_id=getattr(current_user, "id", None),
        )
        return "organiser"

    # If neither token nor user is present or matches
    raise HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Authentication required. Provide a valid token or session.",
    )


@router.get("", response_model=SpeakerProfileResponse)
async def get_speaker_profile(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    access_type: str = Depends(check_profile_access),
) -> SpeakerProfileResponse:
    """
    Returns SpeakerProfileResponse or 404 if no profile yet.
    """
    profile = await SpeakerQueryService(db).get_profile_for_event(
        organization_id=None,
        event_id=event_id,
        speaker_id=speaker_id,
    )
    if not profile:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Speaker profile not found."
        )
    return profile


@router.put("", response_model=SpeakerProfileResponse)
async def upsert_speaker_profile(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    payload: SpeakerProfileUpdate,
    idempotency_key: Optional[str] = Header(None, alias="Idempotency-Key", min_length=8, max_length=200),
    db: AsyncSession = Depends(get_db),
    access_type: str = Depends(check_profile_access),
) -> SpeakerProfileResponse:
    """
    Creates or updates the speaker profile.
    Sets last_updated_by='organiser' if called by organiser JWT, 'speaker' if called by portal token.
    """
    # Resolve event context once; the command owns profile persistence.
    event = await db.get(Event, event_id)
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Speaker not found.")
    return await SpeakerCommandService.upsert_profile(
        db,
        event=event,
        speaker_id=speaker_id,
        payload=payload,
        access_type=access_type,
        actor_user_id=None,
        idempotency_key=idempotency_key,
    )


def extract_text_from_file(filename: str, content: bytes) -> str:
    ext = filename.lower().split('.')[-1]
    if ext == 'pdf':
        try:
            import pdfplumber
        except ImportError:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="CV/Template PDF parsing not available"
            )
        try:
            text = ""
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to process PDF content: {str(e)}"
            )
    elif ext == 'docx':
        try:
            import docx
        except ImportError:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="python-docx not installed"
            )
        try:
            doc = docx.Document(io.BytesIO(content))
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read DOCX template: {str(e)}"
            )
    elif ext == 'pptx':
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            text_runs = []
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                slide_files = [name for name in z.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
                def get_slide_num(name):
                    match = re.search(r'slide(\d+)\.xml', name)
                    return int(match.group(1)) if match else 0
                slide_files.sort(key=get_slide_num)
                
                for slide_file in slide_files:
                    slide_xml = z.read(slide_file)
                    root = ET.fromstring(slide_xml)
                    for elem in root.iter():
                        if elem.tag.endswith('}t'):
                            if elem.text:
                                text_runs.append(elem.text)
            return "\n".join(text_runs)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to process PPTX content: {str(e)}"
            )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload PDF, DOCX, or PPTX."
        )


@router.post("/parse-cv", response_model=SpeakerProfileUpdate)
async def parse_cv(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    file: UploadFile = File(...),
    idempotency_key: str = Header(
        ..., alias="Idempotency-Key", min_length=8, max_length=200
    ),
    db: AsyncSession = Depends(get_db),
    access_type: str = Depends(check_profile_access),
) -> SpeakerProfileUpdate:
    """
    Directly upload the CV to storage without parsing and save cv_url.
    """
    # Check format
    if not file.filename.lower().endswith((".pdf", ".docx", ".pptx")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF, DOCX, and PPTX CV uploads are supported."
        )

    # Check size (max 5MB)
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File size exceeds 5MB limit."
        )

    speaker_res = await db.execute(
        select(Speaker).where(Speaker.id == speaker_id, Speaker.event_id == event_id).options(selectinload(Speaker.event))
    )
    speaker = speaker_res.scalar_one_or_none()
    if not speaker:
        raise HTTPException(status_code=404, detail="Speaker not found.")
    reservation = await UsageReservationService.reserve(
        db,
        organization_id=speaker.event.organization_id,
        event_id=event_id,
        limit_key="storage_quota_mb",
        quantity=max(1, (len(content) + 1024 * 1024 - 1) // (1024 * 1024)),
        unit="megabyte",
        idempotency_key=f"speaker-profile-cv:{speaker_id}:{idempotency_key}",
        metadata={
            "speaker_id": str(speaker_id),
            "profile_asset": "cv",
            "consumption_quantity": len(content),
            "consumption_unit": "byte",
        },
    )

    ext = file.filename.split('.')[-1].lower()
    from app.services import upload_service
    from app.config import settings

    storage_path = f"{speaker.event.organization_id}/{speaker.event_id}/speakers/{speaker.id}/cv.{ext}"
    bucket = settings.S3_BUCKET_ASSETS

    try:
        await asyncio.to_thread(
            upload_service.upload_bytes,
            bucket=bucket,
            storage_path=storage_path,
            data=content,
            content_type=file.content_type,
        )
        if settings.STORAGE_MODE == "local":
            url = f"{settings.API_BASE_URL}{settings.api_v1_prefix}/storage/{bucket}/{storage_path}"
        else:
            url = f"{settings.S3_ENDPOINT_URL}/{bucket}/{storage_path}"
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to upload CV to storage: {str(e)}"
        )

    profile_res = await db.execute(
        select(SpeakerProfile).where(
            SpeakerProfile.speaker_id == speaker_id,
            SpeakerProfile.event_id == event_id
        )
    )
    profile = profile_res.scalar_one_or_none()
    if not profile:
        profile = SpeakerProfile(
            speaker_id=speaker_id,
            event_id=event_id,
            organization_id=speaker.event.organization_id,
            cv_url=url,
            last_updated_by=access_type
        )
        db.add(profile)
    else:
        profile.cv_url = url
        profile.last_updated_by = access_type

    await UsageReservationService.consume(
        db,
        reservation.id,
        source="speaker_profile.cv_upload",
    )
    await db.commit()
    await db.refresh(profile)
    return SpeakerProfileUpdate(cv_url=url)


@router.get("/template")
async def get_profile_template(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    access_type: str = Depends(check_profile_access),
):
    """
    Generates a DOCX file using python-docx pre-filled with speaker and talk info.
    """
    try:
        import docx
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="python-docx not installed"
        )

    # 1. Fetch speaker
    speaker_res = await db.execute(
        select(Speaker).where(Speaker.id == speaker_id, Speaker.event_id == event_id)
    )
    speaker = speaker_res.scalar_one_or_none()
    if not speaker:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Speaker not found.")

    # 2. Fetch event
    event = await db.get(Event, event_id)
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Event not found.")

    # Redirect to custom template if uploaded
    speaker_settings = event.speaker_settings or {}
    profile_settings = speaker_settings.get("profile_settings", {})
    custom_template_url = profile_settings.get("template_url")
    if custom_template_url:
        from fastapi.responses import RedirectResponse
        return RedirectResponse(url=custom_template_url)

    # 3. Fetch session speakers
    ss_result = await db.execute(
        select(SessionSpeaker)
        .join(Session, SessionSpeaker.session_id == Session.id)
        .where(SessionSpeaker.speaker_id == speaker_id, Session.event_id == event_id)
        .options(selectinload(SessionSpeaker.session))
    )
    session_speakers = ss_result.scalars().all()

    # Create DOCX
    doc = docx.Document()
    doc.add_heading("Speaker Profile Template", level=0)

    doc.add_paragraph(f"Speaker Name: {speaker.full_name}")
    doc.add_paragraph(f"Event Name: {event.name}")
    doc.add_paragraph(f"Event Dates: {event.start_date} to {event.end_date}")

    doc.add_heading("Assigned Talks", level=2)
    if not session_speakers:
        doc.add_paragraph("No oral presentation talks assigned yet.")
    else:
        for ss in session_speakers:
            doc.add_paragraph(f"- {ss.presentation_title or 'Untitled Talk'} (Session: {ss.session.name})")

    doc.add_heading("Speaker Profile Details", level=2)
    doc.add_paragraph("Please enter your details below. Do not delete or rename the bracketed section labels.")

    sections = [
        ("[DESIGNATION]", "e.g., Dr. / Prof. / Mr. / Ms."),
        ("[SHORT BIO]", "A short bio (max 150 words)"),
        ("[FULL BIO]", "A full extended bio (max 400 words)"),
        ("[ORGANISATION]", "Your university or company name"),
        ("[DEPARTMENT]", "Your department name"),
        ("[RESEARCH INTERESTS — comma separated]", "e.g., Artificial Intelligence, Genomics"),
        ("[WEBSITE]", "e.g., https://academic-page.edu"),
        ("[LINKEDIN]", "e.g., https://linkedin.com/in/username"),
    ]

    for label, placeholder in sections:
        doc.add_paragraph(f"{label}\n{placeholder}\n")

    # Save to buffer
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    filename = f"{speaker.full_name.replace(' ', '_')}_profile_template.docx"
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/parse-template", response_model=SpeakerProfileUpdate)
async def parse_profile_template(
    event_id: uuid.UUID,
    speaker_id: uuid.UUID,
    file: UploadFile = File(...),
    idempotency_key: str = Header(
        ..., alias="Idempotency-Key", min_length=8, max_length=200
    ),
    db: AsyncSession = Depends(get_db),
    access_type: str = Depends(check_profile_access),
) -> SpeakerProfileUpdate:
    """
    Directly upload the completed template to storage without parsing and save template_url.
    """
    # Check format
    if not file.filename.lower().endswith((".pdf", ".docx", ".pptx")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF, DOCX, and PPTX template uploads are supported."
        )

    # Check size (max 5MB)
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File size exceeds 5MB limit."
        )

    speaker_res = await db.execute(
        select(Speaker).where(Speaker.id == speaker_id, Speaker.event_id == event_id).options(selectinload(Speaker.event))
    )
    speaker = speaker_res.scalar_one_or_none()
    if not speaker:
        raise HTTPException(status_code=404, detail="Speaker not found.")
    reservation = await UsageReservationService.reserve(
        db,
        organization_id=speaker.event.organization_id,
        event_id=event_id,
        limit_key="storage_quota_mb",
        quantity=max(1, (len(content) + 1024 * 1024 - 1) // (1024 * 1024)),
        unit="megabyte",
        idempotency_key=f"speaker-profile-template:{speaker_id}:{idempotency_key}",
        metadata={
            "speaker_id": str(speaker_id),
            "profile_asset": "template",
            "consumption_quantity": len(content),
            "consumption_unit": "byte",
        },
    )

    ext = file.filename.split('.')[-1].lower()
    from app.services import upload_service
    from app.config import settings

    storage_path = f"{speaker.event.organization_id}/{speaker.event_id}/speakers/{speaker.id}/template.{ext}"
    bucket = settings.S3_BUCKET_ASSETS

    try:
        await asyncio.to_thread(
            upload_service.upload_bytes,
            bucket=bucket,
            storage_path=storage_path,
            data=content,
            content_type=file.content_type,
        )
        if settings.STORAGE_MODE == "local":
            url = f"{settings.API_BASE_URL}{settings.api_v1_prefix}/storage/{bucket}/{storage_path}"
        else:
            url = f"{settings.S3_ENDPOINT_URL}/{bucket}/{storage_path}"
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to upload template to storage: {str(e)}"
        )

    profile_res = await db.execute(
        select(SpeakerProfile).where(
            SpeakerProfile.speaker_id == speaker_id,
            SpeakerProfile.event_id == event_id
        )
    )
    profile = profile_res.scalar_one_or_none()
    if not profile:
        profile = SpeakerProfile(
            speaker_id=speaker_id,
            event_id=event_id,
            organization_id=speaker.event.organization_id,
            template_url=url,
            last_updated_by=access_type
        )
        db.add(profile)
    else:
        profile.template_url = url
        profile.last_updated_by = access_type

    await UsageReservationService.consume(
        db,
        reservation.id,
        source="speaker_profile.template_upload",
    )
    await db.commit()
    await db.refresh(profile)
    return SpeakerProfileUpdate(template_url=url)
