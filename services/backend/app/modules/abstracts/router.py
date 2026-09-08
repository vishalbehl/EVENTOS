from __future__ import annotations

import io
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Literal, Optional

from fastapi import APIRouter, Body, Depends, Header, HTTPException, Query, Response, status
from pydantic import BaseModel, ConfigDict, EmailStr, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies.feature_gate import require_event_operation
from app.dependencies import CurrentEvent, get_current_user, get_db
from app.modules.abstracts.models import (
    AbstractAssignment,
    AbstractAttachment,
    AbstractAuthor,
    AbstractCall,
    AbstractDecision,
    AbstractForm,
    AbstractReview,
    AbstractReviewer,
    AbstractSubmission,
)
from app.modules.abstracts.application.commands import AbstractConfigurationCommandService
from app.modules.abstracts.application.submission_commands import AbstractSubmissionCommandService
from app.modules.abstracts.application.reviewer_commands import AbstractReviewerCommandService
from app.modules.abstracts.application.queries import AbstractQueryService
from app.modules.abstracts.application.assignment_commands import AbstractAssignmentCommandService
from app.modules.abstracts.application.decision_commands import AbstractDecisionCommandService
from app.modules.abstracts.application.publication_commands import AbstractPublicationCommandService
from app.modules.abstracts.application.review_commands import AbstractReviewCommandService
from app.modules.abstracts.application.attachment_commands import AbstractAttachmentCommandService
from app.modules.agenda.models import Session
from app.modules.events.models.speaker import Speaker
from app.modules.identity.models.user import User

router = APIRouter(prefix="/events/{event_id}/abstracts", tags=["abstracts"])


class AbstractCallPayload(BaseModel):
    status: Literal["DRAFT", "OPEN", "CLOSED", "PUBLISHED"] = "DRAFT"
    opens_at: Optional[datetime] = None
    closes_at: Optional[datetime] = None
    revision_deadline: Optional[datetime] = None
    min_words: int = Field(default=50, ge=1, le=2000)
    max_words: int = Field(default=300, ge=1, le=5000)
    abstract_types: list[str] = Field(default_factory=lambda: ["ORAL", "POSTER"])
    topics: list[str] = Field(default_factory=list)
    author_rules: dict[str, Any] = Field(default_factory=dict)
    attachment_rules: dict[str, Any] = Field(default_factory=dict)
    disclosure_rules: dict[str, Any] = Field(default_factory=dict)
    email_triggers: dict[str, Any] = Field(default_factory=dict)
    blind_review_enabled: bool = False


class AbstractCallResponse(AbstractCallPayload):
    id: uuid.UUID
    event_id: uuid.UUID
    version: int
    published_at: Optional[datetime]


class AbstractFormPayload(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    title: str = Field(default="Abstract submission form", min_length=3, max_length=180)
    form_schema: dict[str, Any] = Field(default_factory=dict, alias="schema")
    is_active: bool = True
    publish: bool = False


class AbstractFormResponse(AbstractFormPayload):
    id: uuid.UUID
    event_id: uuid.UUID
    version: int
    published_at: Optional[datetime]


class AbstractSubmissionPayload(BaseModel):
    title: str = Field(min_length=3, max_length=300)
    body: str = Field(min_length=10, max_length=20000)
    keywords: list[str] = Field(default_factory=list, max_length=20)
    abstract_type: str = Field(default="ORAL", max_length=40)
    topic: Optional[str] = Field(None, max_length=120)
    track: Optional[str] = Field(None, max_length=120)
    presenter_speaker_id: Optional[uuid.UUID] = None
    submitter_speaker_id: Optional[uuid.UUID] = None
    linked_session_id: Optional[uuid.UUID] = None
    form_payload: dict[str, Any] = Field(default_factory=dict)
    authors: list[dict[str, Any]] = Field(default_factory=list)


class AbstractSubmissionResponse(BaseModel):
    id: uuid.UUID
    event_id: uuid.UUID
    code: str
    title: str
    body: str
    keywords: list[str]
    abstract_type: str
    topic: Optional[str]
    track: Optional[str]
    status: str
    version: int
    submitted_at: Optional[datetime]
    decided_at: Optional[datetime]
    published_at: Optional[datetime]
    presenter_name: Optional[str] = None
    presenter_email: Optional[str] = None
    linked_session_name: Optional[str] = None
    average_score: Optional[int] = None
    review_count: int = 0
    assignment_count: int = 0
    conflict_count: int = 0
    final_decision: Optional[str] = None
    presentation_type: Optional[str] = None
    authors: list[dict[str, Any]] = Field(default_factory=list)
    attachments: list[dict[str, Any]] = Field(default_factory=list)
    reviews: list[dict[str, Any]] = Field(default_factory=list)
    decisions: list[dict[str, Any]] = Field(default_factory=list)


class AbstractDashboardResponse(BaseModel):
    call: Optional[AbstractCallResponse]
    counts: dict[str, int]
    reviewer_load: dict[str, int]
    publication: dict[str, int]
    needs_attention: list[dict[str, str]]
    recent: list[AbstractSubmissionResponse]


class AbstractSubmissionPage(BaseModel):
    items: list[AbstractSubmissionResponse]
    next_cursor: Optional[uuid.UUID] = None


class ReviewerPayload(BaseModel):
    full_name: str = Field(min_length=2, max_length=180)
    email: EmailStr
    expertise_topics: list[str] = Field(default_factory=list)
    capacity: int = Field(default=10, ge=1, le=100)
    status: Literal["INVITED", "ACTIVE", "PAUSED"] = "INVITED"


class ReviewerResponse(ReviewerPayload):
    id: uuid.UUID
    assigned_count: int = 0
    completed_count: int = 0
    conflict_count: int = 0


class ReviewerUpdatePayload(BaseModel):
    full_name: Optional[str] = Field(None, min_length=2, max_length=180)
    email: Optional[EmailStr] = None
    expertise_topics: Optional[list[str]] = None
    capacity: Optional[int] = Field(None, ge=1, le=100)
    status: Optional[Literal["INVITED", "ACTIVE", "PAUSED"]] = None


class AssignmentPayload(BaseModel):
    reviewer_id: uuid.UUID
    due_at: Optional[datetime] = None


class AssignmentUpdatePayload(BaseModel):
    status: Optional[Literal["ASSIGNED", "ACCEPTED", "DECLINED", "CONFLICT", "COMPLETED"]] = None
    due_at: Optional[datetime] = None
    conflict_declared: Optional[bool] = None
    conflict_reason: Optional[str] = Field(None, max_length=2000)


class ReviewPayload(BaseModel):
    scores: dict[str, int] = Field(default_factory=dict)
    recommendation: Literal["ACCEPT", "REJECT", "REVISION", "DISCUSS"] = "DISCUSS"
    comments_to_committee: Optional[str] = Field(None, max_length=4000)
    comments_to_author: Optional[str] = Field(None, max_length=4000)


class DecisionPayload(BaseModel):
    decision: Literal["ACCEPTED", "REJECTED", "REVISION_REQUESTED"]
    presentation_type: Optional[Literal["ORAL", "POSTER", "EPOSTER"]] = None
    reason: str = Field(min_length=5, max_length=1000)
    notes_to_author: Optional[str] = Field(None, max_length=4000)


class PublicationPayload(BaseModel):
    publication_payload: dict[str, Any] = Field(default_factory=dict)


class AttachmentPayload(BaseModel):
    kind: str = Field(default="SUPPORTING_FILE", max_length=40)
    filename: str = Field(min_length=1, max_length=255)
    storage_path: Optional[str] = Field(None, max_length=600)
    mime_type: Optional[str] = Field(None, max_length=120)
    file_size_bytes: Optional[int] = Field(None, ge=0)


def _call_response(row: AbstractCall) -> AbstractCallResponse:
    return AbstractCallResponse(**{field: getattr(row, field) for field in AbstractCallResponse.model_fields})


def _form_response(row: AbstractForm) -> AbstractFormResponse:
    return AbstractFormResponse(
        id=row.id,
        event_id=row.event_id,
        title=row.title,
        schema=row.schema or {},
        is_active=row.is_active,
        publish=bool(row.published_at),
        version=row.version,
        published_at=row.published_at,
    )


async def _submission_response(db: AsyncSession, row: AbstractSubmission) -> AbstractSubmissionResponse:
    return (await _submission_responses(db, [row], row.event_id))[0]


def _build_submission_response(
    row: AbstractSubmission,
    *,
    speaker_by_id: dict[uuid.UUID, Speaker],
    session_by_id: dict[uuid.UUID, Session],
    assignment_stats: dict[uuid.UUID, tuple[int, int]],
    review_counts: dict[uuid.UUID, int],
    authors_by_submission: dict[uuid.UUID, list[AbstractAuthor]],
    attachments_by_submission: dict[uuid.UUID, list[AbstractAttachment]],
    reviews_by_submission: dict[uuid.UUID, list[tuple[AbstractReview, AbstractReviewer]]],
    decisions_by_submission: dict[uuid.UUID, list[AbstractDecision]],
) -> AbstractSubmissionResponse:
    speaker = speaker_by_id.get(row.presenter_speaker_id or row.submitter_speaker_id)
    session = session_by_id.get(row.linked_session_id)
    authors = authors_by_submission.get(row.id, [])
    attachments = attachments_by_submission.get(row.id, [])
    reviews = reviews_by_submission.get(row.id, [])
    decisions = decisions_by_submission.get(row.id, [])
    assignments, conflicts = assignment_stats.get(row.id, (0, 0))
    return AbstractSubmissionResponse(
        id=row.id,
        event_id=row.event_id,
        code=row.code,
        title=row.title,
        body=row.body,
        keywords=row.keywords or [],
        abstract_type=row.abstract_type,
        topic=row.topic,
        track=row.track,
        status=row.status,
        version=row.version,
        submitted_at=row.submitted_at,
        decided_at=row.decided_at,
        published_at=row.published_at,
        presenter_name=speaker.full_name if speaker else None,
        presenter_email=speaker.email if speaker else None,
        linked_session_name=session.name if session else None,
        average_score=row.average_score,
        review_count=review_counts.get(row.id, 0),
        assignment_count=assignments,
        conflict_count=conflicts,
        final_decision=row.final_decision,
        presentation_type=row.presentation_type,
        authors=[
            {
                "id": str(author.id),
                "full_name": author.full_name,
                "email": author.email,
                "affiliation": author.affiliation,
                "country": author.country,
                "is_presenter": author.is_presenter,
                "display_order": author.display_order,
            }
            for author in authors
        ],
        attachments=[
            {
                "id": str(attachment.id),
                "kind": attachment.kind,
                "filename": attachment.filename,
                "storage_path": attachment.storage_path,
                "mime_type": attachment.mime_type,
                "file_size_bytes": attachment.file_size_bytes,
            }
            for attachment in attachments
        ],
        reviews=[
            {
                "id": str(review.id),
                "reviewer_id": str(reviewer.id),
                "reviewer_name": reviewer.full_name,
                "scores": review.scores,
                "total_score": review.total_score,
                "recommendation": review.recommendation,
                "comments_to_committee": review.comments_to_committee,
                "comments_to_author": review.comments_to_author,
                "submitted_at": review.submitted_at.isoformat(),
            }
            for review, reviewer in reviews
        ],
        decisions=[
            {
                "id": str(decision.id),
                "decision": decision.decision,
                "presentation_type": decision.presentation_type,
                "reason": decision.reason,
                "notes_to_author": decision.notes_to_author,
                "decided_at": decision.decided_at.isoformat(),
            }
            for decision in decisions
        ],
    )


async def _submission_responses(
    db: AsyncSession,
    rows: list[AbstractSubmission],
    event_id: uuid.UUID,
) -> list[AbstractSubmissionResponse]:
    """Build submission projections with bounded batched related reads."""
    if not rows:
        return []
    submission_ids = [row.id for row in rows]
    speaker_ids = {
        speaker_id
        for row in rows
        for speaker_id in (row.presenter_speaker_id, row.submitter_speaker_id)
        if speaker_id
    }
    session_ids = {row.linked_session_id for row in rows if row.linked_session_id}
    related = await AbstractQueryService(db).submission_related(
        event_id=event_id, submission_ids=submission_ids
    )
    speakers = related.speakers
    sessions = related.sessions
    assignment_rows = related.assignment_rows
    review_rows = related.review_rows
    authors = related.authors
    attachments = related.attachments
    reviews = related.reviews
    decisions = related.decisions

    assignment_stats = {item[0]: (int(item[1] or 0), int(item[2] or 0)) for item in assignment_rows}
    review_counts = {item[0]: int(item[1] or 0) for item in review_rows}
    authors_by_submission = defaultdict(list)
    for author in authors:
        authors_by_submission[author.submission_id].append(author)
    attachments_by_submission = defaultdict(list)
    for attachment in attachments:
        attachments_by_submission[attachment.submission_id].append(attachment)
    reviews_by_submission = defaultdict(list)
    for review, reviewer in reviews:
        reviews_by_submission[review.submission_id].append((review, reviewer))
    decisions_by_submission = defaultdict(list)
    for decision in decisions:
        decisions_by_submission[decision.submission_id].append(decision)
    speaker_by_id = {speaker.id: speaker for speaker in speakers}
    session_by_id = {session.id: session for session in sessions}
    return [
        _build_submission_response(
            row,
            speaker_by_id=speaker_by_id,
            session_by_id=session_by_id,
            assignment_stats=assignment_stats,
            review_counts=review_counts,
            authors_by_submission=authors_by_submission,
            attachments_by_submission=attachments_by_submission,
            reviews_by_submission=reviews_by_submission,
            decisions_by_submission=decisions_by_submission,
        )
        for row in rows
    ]


async def _get_submission(db: AsyncSession, event: CurrentEvent, submission_id: uuid.UUID) -> AbstractSubmission:
    row = await AbstractQueryService(db).get_submission(
        event_id=event.id, submission_id=submission_id
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Submission not found")
    return row


@router.get("/dashboard", response_model=AbstractDashboardResponse, dependencies=[require_event_operation("abstracts.review")])
async def abstract_dashboard(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> AbstractDashboardResponse:
    dashboard = await AbstractQueryService(db).dashboard(event_id=event.id)
    call = dashboard.call
    counts = dashboard.submission_counts
    total_reviewers = dashboard.reviewer_count
    assigned = dashboard.assigned_count
    completed = dashboard.completed_review_count
    accepted = int(counts.get("ACCEPTED", 0))
    published = dashboard.published_count
    needs_attention: list[dict[str, str]] = []
    if call is None or call.status == "DRAFT":
        needs_attention.append({"severity": "warning", "label": "Call for abstracts is not open", "destination": f"/events/{event.id}/abstracts/setup"})
    if counts.get("SUBMITTED", 0) and total_reviewers == 0:
        needs_attention.append({"severity": "critical", "label": "Submissions are waiting without reviewers", "destination": f"/events/{event.id}/abstracts/reviewers"})
    if accepted > published:
        needs_attention.append({"severity": "warning", "label": "Accepted abstracts are not fully published", "destination": f"/events/{event.id}/abstracts/accepted"})
    return AbstractDashboardResponse(
        call=_call_response(call) if call else None,
        counts=counts,
        reviewer_load={"reviewers": total_reviewers, "assigned": assigned, "completed": completed},
        publication={"accepted": accepted, "published": published, "ready": max(accepted - published, 0)},
        needs_attention=needs_attention,
        recent=await _submission_responses(db, dashboard.recent_submissions, event.id),
    )


@router.get("/setup", response_model=AbstractCallResponse, dependencies=[require_event_operation("abstracts.configure")])
async def get_setup(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> AbstractCallResponse:
    call = await AbstractQueryService(db).get_call(event_id=event.id)
    if call is None:
        # GET must remain read-only. The first explicit PUT creates the durable
        # setup record; this projection keeps the existing response contract.
        call = AbstractCall(
            id=uuid.uuid4(),
            organization_id=event.organization_id,
            event_id=event.id,
            version=1,
            status="DRAFT",
            max_words=300,
            min_words=50,
            abstract_types=["ORAL", "POSTER"],
            topics=[],
            author_rules={},
            attachment_rules={},
            disclosure_rules={},
            email_triggers={},
            blind_review_enabled=False,
            published_at=None,
        )
    return _call_response(call)


@router.put("/setup", response_model=AbstractCallResponse, dependencies=[require_event_operation("abstracts.configure")])
async def update_setup(payload: AbstractCallPayload, event: CurrentEvent, expected_version: Optional[int] = Header(None, alias="If-Match"), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractCallResponse:
    call = await AbstractConfigurationCommandService(db).update_call(payload=payload.model_dump(), event=event, user=current_user, expected_version=expected_version)
    return _call_response(call)


@router.get("/form", response_model=AbstractFormResponse, dependencies=[require_event_operation("abstracts.configure")])
async def get_form(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> AbstractFormResponse:
    form = await AbstractQueryService(db).get_active_form(event_id=event.id)
    if form is None:
        # Return the same default form without seeding it from a read request.
        form = AbstractForm(
            id=uuid.uuid4(),
            organization_id=event.organization_id,
            event_id=event.id,
            title="Abstract submission form",
            schema={"fields": ["title", "body", "keywords", "topic", "authors"]},
            version=1,
            is_active=True,
            published_at=None,
        )
    return _form_response(form)


@router.put("/form", response_model=AbstractFormResponse, dependencies=[require_event_operation("abstracts.configure")])
async def update_form(payload: AbstractFormPayload, event: CurrentEvent, expected_version: Optional[int] = Header(None, alias="If-Match"), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractFormResponse:
    form = await AbstractConfigurationCommandService(db).update_form(payload=payload.model_dump(), event=event, user=current_user, expected_version=expected_version)
    return _form_response(form)


@router.get("/submissions", response_model=AbstractSubmissionPage, dependencies=[require_event_operation("abstracts.review")])
async def list_submissions(event: CurrentEvent, status_filter: Optional[str] = Query(None, alias="status"), search: Optional[str] = Query(None, max_length=120), topic: Optional[str] = None, cursor: Optional[uuid.UUID] = None, limit: int = Query(100, ge=1, le=200), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionPage:
    page_rows, has_more = await AbstractQueryService(db).list_submissions(
        event_id=event.id,
        status_filter=status_filter,
        search=search,
        topic=topic,
        cursor=cursor,
        limit=limit,
    )
    return AbstractSubmissionPage(
        items=await _submission_responses(db, page_rows, event.id),
        next_cursor=page_rows[-1].id if has_more and page_rows else None,
    )


@router.post("/submissions", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.submit")])
async def create_submission(payload: AbstractSubmissionPayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    row = await AbstractSubmissionCommandService(db).create(
        payload=payload.model_dump(), event=event, user=current_user
    )
    return await _submission_response(db, row)


@router.get("/submissions/{submission_id}", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.review")])
async def get_submission(submission_id: uuid.UUID, event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    return await _submission_response(db, await _get_submission(db, event, submission_id))


@router.put("/submissions/{submission_id}", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.review")])
async def update_submission(submission_id: uuid.UUID, payload: AbstractSubmissionPayload, event: CurrentEvent, expected_version: int = Header(..., alias="If-Match", ge=1), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    row = await AbstractSubmissionCommandService(db).update(
        submission_id=submission_id,
        payload=payload.model_dump(),
        expected_version=expected_version,
        event=event,
        user=current_user,
    )
    return await _submission_response(db, row)


@router.delete("/submissions/{submission_id}", status_code=status.HTTP_204_NO_CONTENT, dependencies=[require_event_operation("abstracts.review")])
async def delete_submission(submission_id: uuid.UUID, event: CurrentEvent, expected_version: int = Header(..., alias="If-Match", ge=1), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> Response:
    await AbstractSubmissionCommandService(db).delete(
        submission_id=submission_id,
        expected_version=expected_version,
        event=event,
        user=current_user,
    )
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/submissions/{submission_id}/attachments", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.review")])
async def add_attachment(submission_id: uuid.UUID, payload: AttachmentPayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    row = await AbstractAttachmentCommandService(db).add(
        submission_id=submission_id,
        payload=payload.model_dump(),
        event=event,
        user=current_user,
    )
    return await _submission_response(db, row)


@router.get("/reviewers", response_model=list[ReviewerResponse], dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def list_reviewers(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> list[ReviewerResponse]:
    rows = await AbstractQueryService(db).list_reviewers(event_id=event.id)
    return [
        ReviewerResponse(
            id=reviewer.id,
            full_name=reviewer.full_name,
            email=reviewer.email,
            expertise_topics=reviewer.expertise_topics or [],
            capacity=reviewer.capacity,
            status=reviewer.status,
            assigned_count=row.assigned_count,
            completed_count=row.completed_count,
            conflict_count=row.conflict_count,
        )
        for row in rows
        for reviewer in (row.reviewer,)
    ]


@router.post("/reviewers", response_model=ReviewerResponse, dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def create_reviewer(payload: ReviewerPayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> ReviewerResponse:
    reviewer = await AbstractReviewerCommandService(db).create(payload=payload.model_dump(), event=event, user=current_user)
    return ReviewerResponse(id=reviewer.id, full_name=reviewer.full_name, email=reviewer.email, expertise_topics=reviewer.expertise_topics or [], capacity=reviewer.capacity, status=reviewer.status)


@router.put("/reviewers/{reviewer_id}", response_model=ReviewerResponse, dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def update_reviewer(reviewer_id: uuid.UUID, payload: ReviewerUpdatePayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> ReviewerResponse:
    reviewer = await AbstractReviewerCommandService(db).update(
        reviewer_id=reviewer_id, payload=payload.model_dump(exclude_unset=True), event=event, user=current_user
    )
    assigned, completed, conflicts = await AbstractQueryService(db).reviewer_stats(
        event_id=event.id, reviewer_id=reviewer.id
    )
    return ReviewerResponse(id=reviewer.id, full_name=reviewer.full_name, email=reviewer.email, expertise_topics=reviewer.expertise_topics or [], capacity=reviewer.capacity, status=reviewer.status, assigned_count=assigned, completed_count=completed, conflict_count=conflicts)


@router.delete("/reviewers/{reviewer_id}", status_code=status.HTTP_204_NO_CONTENT, dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def delete_reviewer(reviewer_id: uuid.UUID, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> Response:
    await AbstractReviewerCommandService(db).delete(reviewer_id=reviewer_id, event=event, user=current_user)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/assignments", dependencies=[require_event_operation("abstracts.review")])
async def list_assignments(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> list[dict[str, Any]]:
    rows = await AbstractQueryService(db).list_assignments(event_id=event.id)
    return [
        {
            "id": str(row.assignment.id),
            "submission_id": str(row.submission.id),
            "submission_code": row.submission.code,
            "submission_title": row.submission.title,
            "reviewer_id": str(row.reviewer.id),
            "reviewer_name": row.reviewer.full_name,
            "status": row.assignment.status,
            "due_at": row.assignment.due_at.isoformat() if row.assignment.due_at else None,
            "conflict_declared": row.assignment.conflict_declared,
            "conflict_reason": row.assignment.conflict_reason,
        }
        for row in rows
    ]


@router.post("/submissions/{submission_id}/assignments", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def assign_reviewer(submission_id: uuid.UUID, payload: AssignmentPayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    submission = await AbstractAssignmentCommandService(db).assign(
        submission_id=submission_id, reviewer_id=payload.reviewer_id, due_at=payload.due_at,
        event=event, user=current_user
    )
    return await _submission_response(db, submission)


@router.put("/assignments/{assignment_id}", dependencies=[require_event_operation("abstracts.assign_reviewers")])
async def update_assignment(assignment_id: uuid.UUID, payload: AssignmentUpdatePayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> dict[str, Any]:
    return await AbstractAssignmentCommandService(db).update(
        assignment_id=assignment_id, payload=payload.model_dump(exclude_unset=True), event=event, user=current_user
    )


@router.post("/assignments/{assignment_id}/reviews", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.review")])
async def submit_review(assignment_id: uuid.UUID, payload: ReviewPayload, event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    submission = await AbstractReviewCommandService(db).submit(
        assignment_id=assignment_id,
        payload=payload.model_dump(),
        event=event,
        user=current_user,
    )
    return await _submission_response(db, submission)


@router.post("/submissions/{submission_id}/decision", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.decide")])
async def decide_submission(submission_id: uuid.UUID, payload: DecisionPayload, event: CurrentEvent, expected_version: int = Header(..., alias="If-Match", ge=1), idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=8, max_length=200), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    submission = await AbstractDecisionCommandService(db).decide(
        submission_id=submission_id,
        decision=payload.decision,
        presentation_type=payload.presentation_type,
        reason=payload.reason,
        notes_to_author=payload.notes_to_author,
        expected_version=expected_version,
        idempotency_key=idempotency_key,
        event=event,
        user=current_user,
    )
    return await _submission_response(db, submission)


@router.post("/submissions/{submission_id}/publish", response_model=AbstractSubmissionResponse, dependencies=[require_event_operation("abstracts.publish")])
async def publish_submission(submission_id: uuid.UUID, event: CurrentEvent, payload: PublicationPayload | None = Body(None), current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionResponse:
    submission = await AbstractPublicationCommandService(db).publish(
        submission_id=submission_id,
        publication_payload=payload.publication_payload if payload else {},
        event=event,
        user=current_user,
    )
    return await _submission_response(db, submission)


@router.post("/accepted/publish-all", response_model=AbstractSubmissionPage, dependencies=[require_event_operation("abstracts.publish")])
async def publish_all_accepted(event: CurrentEvent, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)) -> AbstractSubmissionPage:
    await AbstractPublicationCommandService(db).publish_all(event=event, user=current_user)
    refreshed = await AbstractQueryService(db).list_accepted_submissions(event_id=event.id)
    return AbstractSubmissionPage(items=await _submission_responses(db, refreshed, event.id))


@router.get("/accepted", response_model=AbstractSubmissionPage, dependencies=[require_event_operation("abstracts.publish")])
async def accepted_directory(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> AbstractSubmissionPage:
    rows = await AbstractQueryService(db).list_accepted_submissions(event_id=event.id)
    return AbstractSubmissionPage(items=await _submission_responses(db, rows, event.id))


@router.get("/exports/manifest", dependencies=[require_event_operation("abstracts.export")])
async def export_manifest(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> dict[str, Any]:
    rows = await AbstractQueryService(db).list_submissions_for_export(event_id=event.id)
    return {
        "event_id": str(event.id),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "items": [
            {
                "code": row.code,
                "title": row.title,
                "status": row.status,
                "type": row.abstract_type,
                "topic": row.topic,
                "track": row.track,
                "score": row.average_score,
                "published": bool(row.published_at),
            }
            for row in rows
        ],
    }


@router.get("/exports/submissions.csv", dependencies=[require_event_operation("abstracts.export")])
async def export_submissions_csv(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> Response:
    rows = await AbstractQueryService(db).list_submissions_for_export(event_id=event.id)
    columns = ["code", "title", "status", "abstract_type", "topic", "track", "average_score", "review_count", "assignment_count", "published"]

    def cell(value: Any) -> str:
        text = "" if value is None else str(value)
        return '"' + text.replace('"', '""') + '"'

    lines = [",".join(columns)]
    for row in rows:
        response = await _submission_response(db, row)
        lines.append(
            ",".join(
                cell(value)
                for value in [
                    response.code,
                    response.title,
                    response.status,
                    response.abstract_type,
                    response.topic,
                    response.track,
                    response.average_score,
                    response.review_count,
                    response.assignment_count,
                    bool(response.published_at),
                ]
            )
        )
    return Response(
        "\n".join(lines),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="event-{event.id}-abstracts.csv"'},
    )


async def _accepted_export_rows(event: CurrentEvent, db: AsyncSession) -> list[AbstractSubmissionResponse]:
    rows = await AbstractQueryService(db).list_submissions_for_export(
        event_id=event.id, accepted_only=True
    )
    return [await _submission_response(db, row) for row in rows]


def _html_escape(value: str) -> str:
    return (value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


@router.get("/exports/abstract-book.html", dependencies=[require_event_operation("abstracts.export")])
async def export_abstract_book_html(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> Response:
    """Return a print-ready abstract book that organisers can save as PDF."""
    rows = await _accepted_export_rows(event, db)
    entries = []
    for row in rows:
        authors = ", ".join(author.get("full_name", "") for author in row.authors if author.get("full_name"))
        entries.append(f"<article><p class='code'>{row.code} · {row.presentation_type or row.abstract_type}</p><h2>{_html_escape(row.title)}</h2><p class='authors'>{_html_escape(authors or row.presenter_name or 'Author not recorded')}</p><p>{_html_escape(row.body)}</p></article>")
    body = "".join(entries) or "<p class='empty'>No accepted abstracts have been published for this event.</p>"
    title = _html_escape(getattr(event, "name", None) or "Abstract book")
    html = f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title><style>@page{{size:A4;margin:18mm}}body{{font:11pt Arial;color:#17202a;line-height:1.55}}h1{{font-size:25pt;margin-bottom:4pt}}h2{{font-size:15pt;margin:4pt 0}}.code,.meta{{color:#5f6b76;font-size:9pt;text-transform:uppercase;letter-spacing:.08em}}article{{break-inside:avoid;border-top:1px solid #d8dde2;padding:14pt 0}}.authors{{font-weight:bold;color:#485563}}.empty{{color:#6b7280}}</style></head><body><p class='meta'>Abstract book</p><h1>{title}</h1><p class='meta'>{len(rows)} accepted abstracts</p>{body}</body></html>"
    return Response(content=html, media_type="text/html", headers={"Content-Disposition": f'inline; filename="event-{event.id}-abstract-book.html"'})


@router.get("/exports/abstract-book.docx", dependencies=[require_event_operation("abstracts.export")])
async def export_abstract_book_docx(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> Response:
    rows = await _accepted_export_rows(event, db)
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="DOCX export dependency is unavailable") from exc
    document = Document()
    document.add_heading(getattr(event, "name", None) or "Abstract book", 0)
    document.add_paragraph(f"{len(rows)} accepted abstracts")
    for row in rows:
        document.add_heading(row.title, 1)
        authors = ", ".join(author.get("full_name", "") for author in row.authors if author.get("full_name"))
        document.add_paragraph(f"{row.code} | {row.presentation_type or row.abstract_type} | {authors or row.presenter_name or 'Author not recorded'}")
        document.add_paragraph(row.body)
    output = io.BytesIO()
    document.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="event-{event.id}-abstract-book.docx"'})


@router.get("/exports/reviewer-report.docx", dependencies=[require_event_operation("abstracts.export")])
async def export_reviewer_report_docx(event: CurrentEvent, db: AsyncSession = Depends(get_db)) -> Response:
    """Export committee evidence as a portable DOCX report."""
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="DOCX export dependency is unavailable") from exc
    rows = await AbstractQueryService(db).list_submissions_for_export(event_id=event.id)
    document = Document()
    document.add_heading(f"Reviewer report - {getattr(event, 'name', None) or 'Event'}", 0)
    document.add_paragraph(f"{len(rows)} submissions | generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    for row in rows:
        response = await _submission_response(db, row)
        document.add_heading(f"{response.code}: {response.title}", 1)
        document.add_paragraph(f"Status: {response.status} | Score: {response.average_score if response.average_score is not None else 'Not scored'} | Reviews: {response.review_count}")
        for review in response.reviews:
            document.add_heading(review.get("reviewer_name", "Reviewer"), 2)
            document.add_paragraph(f"Recommendation: {review.get('recommendation', 'Not recorded')}")
            document.add_paragraph(review.get("comments_to_committee") or "No committee comment recorded.")
    output = io.BytesIO()
    document.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="event-{event.id}-reviewer-report.docx"'})
