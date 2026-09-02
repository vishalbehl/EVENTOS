"""Transaction-owning inventory commands."""

from __future__ import annotations

import io

import openpyxl
from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.identity.models.user import User
from app.modules.inventory.models import HardwareCategory, HardwareItem, HardwareStock


class InventoryCommandService:
    """Import and delete inventory without exposing persistence in the router."""

    MAX_IMPORT_BYTES = 20 * 1024 * 1024
    MAX_IMPORT_ROWS = 10_000

    def __init__(self, db: AsyncSession):
        self.db = db

    @staticmethod
    def _value(row, header: dict[str, int], *keys, default=None):
        for key in keys:
            index = header.get(key)
            if index is not None and index < len(row) and row[index].value is not None:
                return row[index].value
        return default

    async def import_hardware(self, *, contents: bytes, organization_id) -> int:
        if len(contents) > self.MAX_IMPORT_BYTES:
            raise HTTPException(413, "File too large (max 20 MB).")
        try:
            sheet = openpyxl.load_workbook(
                io.BytesIO(contents), data_only=True, read_only=True
            ).active
            header = {
                str(cell.value).strip().lower().replace("_", " "): index
                for index, cell in enumerate(next(sheet.iter_rows()))
                if cell.value
            }
        except Exception as exc:
            raise HTTPException(400, "Invalid Excel file.") from exc

        count = 0
        for row in sheet.iter_rows(min_row=2, max_row=self.MAX_IMPORT_ROWS + 1):
            code = str(
                self._value(row, header, "hardware code", "asset code", "code", default="")
            ).strip()
            if not code:
                continue

            item = await self.db.scalar(
                select(HardwareItem)
                .where(HardwareItem.asset_code == code)
                .execution_options(skip_tenant_filter=True)
            )
            category_name = str(
                self._value(row, header, "category", default="General")
            ).strip()
            category = await self.db.scalar(
                select(HardwareCategory).where(
                    func.lower(HardwareCategory.name) == category_name.lower()
                )
            )
            if category is None:
                category = HardwareCategory(name=category_name)
                self.db.add(category)
                await self.db.flush()

            values = {
                "category_id": category.id,
                "name": str(self._value(row, header, "hardware name", "name", default=code)),
                "brand": str(self._value(row, header, "brand", default="")),
                "model": str(self._value(row, header, "model", default="")),
                "purchase_cost": float(
                    self._value(row, header, "cost price", "purchase cost", default=0) or 0
                ),
                "renting_price": float(
                    self._value(row, header, "renting price", default=0) or 0
                ),
                "pricing_unit": str(
                    self._value(row, header, "pricing unit", default="PER_EVENT")
                ),
                "description": self._value(row, header, "description"),
                "tax_category": self._value(row, header, "tax category"),
            }
            if item is None:
                item = HardwareItem(
                    **values,
                    organization_id=organization_id,
                    asset_code=code,
                )
                self.db.add(item)
                await self.db.flush()
            else:
                for key, value in values.items():
                    setattr(item, key, value)

            quantity = int(
                self._value(row, header, "inventory count", "quantity", default=1) or 1
            )
            stock = await self.db.scalar(
                select(HardwareStock)
                .where(HardwareStock.hardware_id == item.id)
                .with_for_update()
            )
            if stock is None:
                self.db.add(
                    HardwareStock(
                        hardware_id=item.id,
                        quantity=quantity,
                        available_quantity=quantity,
                    )
                )
            else:
                stock.quantity = quantity
                stock.available_quantity = max(0, quantity - stock.reserved_quantity)
            count += 1

        await self.db.commit()
        return count

    async def delete_hardware(self, *, hardware_id) -> None:
        item = await self.db.scalar(
            select(HardwareItem)
            .where(HardwareItem.id == hardware_id)
            .execution_options(skip_tenant_filter=True)
        )
        if item is None:
            raise HTTPException(404, "Hardware item not found")
        await self.db.execute(
            HardwareStock.__table__.delete().where(HardwareStock.hardware_id == hardware_id)
        )
        await self.db.delete(item)
        await self.db.commit()
