import uuid
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user, get_db
from app.modules.identity.models.user import User
from app.modules.inventory.services import InventoryService
from app.modules.inventory.schemas import (
    HardwareItemCreate, HardwareItemOut,
    AllocateHardwareRequest, ReturnHardwareRequest,
    HardwareCategoryCreate, HardwareCategoryOut
)

router = APIRouter(prefix="/inventory", tags=["inventory"])

@router.post("/categories", response_model=HardwareCategoryOut)
async def create_category(
    req: HardwareCategoryCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    try:
        cat = await InventoryService.create_category(
            db=db,
            name=req.name,
            description=req.description
        )
        await db.commit()
        return cat
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=400, detail=str(e))

@router.post("/hardware", response_model=HardwareItemOut)
async def create_hardware(
    req: HardwareItemCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    org_id = current_user.organization_id
    try:
        item = await InventoryService.create_hardware_item(
            db=db,
            organization_id=org_id,
            category_id=req.category_id,
            asset_code=req.asset_code,
            name=req.name,
            brand=req.brand,
            model=req.model,
            serial_number=req.serial_number,
            purchase_date=req.purchase_date,
            purchase_cost=req.purchase_cost,
            replacement_cost=req.replacement_cost,
            status=req.status,
            condition=req.condition,
            location=req.location,
            notes=req.notes
        )
        await db.commit()
        return item
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/hardware", response_model=List[HardwareItemOut])
async def get_hardware(
    status: Optional[str] = None,
    limit: int = 50,
    offset: int = 0,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    org_id = current_user.organization_id
    items = await InventoryService.get_hardware_items(
        db=db,
        organization_id=org_id,
        status=status,
        limit=limit,
        offset=offset
    )
    return items

@router.post("/allocate", response_model=HardwareItemOut)
async def allocate_hardware(
    req: AllocateHardwareRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    org_id = current_user.organization_id
    try:
        item = await InventoryService.allocate_hardware(
            db=db,
            organization_id=org_id,
            hardware_id=req.hardware_id,
            to_location=req.to_location,
            notes=req.notes
        )
        if not item:
            raise HTTPException(status_code=400, detail="Hardware not available for allocation")
        await db.commit()
        return item
    except HTTPException:
        raise
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=400, detail=str(e))

@router.post("/return", response_model=HardwareItemOut)
async def return_hardware(
    req: ReturnHardwareRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    org_id = current_user.organization_id
    try:
        item = await InventoryService.return_hardware(
            db=db,
            organization_id=org_id,
            hardware_id=req.hardware_id,
            return_location=req.return_location,
            condition=req.condition,
            notes=req.notes
        )
        if not item:
            raise HTTPException(status_code=404, detail="Hardware item not found")
        await db.commit()
        return item
    except HTTPException:
        raise
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=400, detail=str(e))


# ── SUPER ADMIN CATALOG ENDPOINTS ──────────────────────────────────────

from pydantic import BaseModel
from sqlalchemy import func, select, text, and_
from datetime import datetime, timezone, timedelta
import json
from app.modules.superadmin.dependencies import require_super_admin
from app.modules.inventory.models import HardwareItem, HardwareCategory, HardwareStock

async def seed_hardware_if_empty(db: AsyncSession):
    # Ensure the 8 default categories exist
    default_categories = [
        "Registration",
        "SRR",
        "Presentation",
        "Networking",
        "Check-in",
        "Office Equipment",
        "Storage",
        "Accessories"
    ]
    
    # Fetch existing categories
    existing_cats = (await db.execute(select(HardwareCategory))).scalars().all()
    existing_names = {c.name for c in existing_cats}
    
    cat_map = {c.name: c.id for c in existing_cats}
    
    # Insert missing categories
    modified = False
    for name in default_categories:
        if name not in existing_names:
            c_id = uuid.uuid4()
            c = HardwareCategory(id=c_id, name=name, description=f"Default {name} hardware category")
            db.add(c)
            cat_map[name] = c_id
            modified = True
            
    if modified:
        await db.flush()
        
    if modified:
        await db.commit()

async def get_next_available_hardware_code(db: AsyncSession, extra_in_use: set = None) -> str:
    # Query all active codes (not retired, not lost, not inactive)
    stmt = select(HardwareItem.asset_code).where(
        HardwareItem.status.notin_(["RETIRED", "LOST", "INACTIVE"])
    )
    res = await db.execute(stmt)
    in_use = {r for r in res.scalars().all() if r}
    if extra_in_use:
        in_use.update(extra_in_use)
        
    num = 1001
    while True:
        candidate = f"HW-{num}"
        if candidate not in in_use:
            return candidate
        num += 1

class SuperAdminHardwareCreate(BaseModel):
    name: str
    item_code: str
    category_id: uuid.UUID
    pricing_unit: str = "PER_EVENT"
    cost_price: float
    selling_price: float
    brand: Optional[str] = "Generic"
    model: Optional[str] = "Generic"
    inventory_count: Optional[int] = 0
    description: Optional[str] = None
    tax_category: Optional[str] = "GST_18"

class SuperAdminHardwareUpdate(BaseModel):
    name: Optional[str] = None
    item_code: Optional[str] = None
    category_id: Optional[uuid.UUID] = None
    pricing_unit: Optional[str] = None
    cost_price: Optional[float] = None
    selling_price: Optional[float] = None
    brand: Optional[str] = None
    model: Optional[str] = None
    inventory_count: Optional[int] = None
    description: Optional[str] = None
    tax_category: Optional[str] = None
    status: Optional[str] = None

@router.get("/superadmin/catalog/hardware")
async def superadmin_get_hardware(
    category: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    pricing_unit: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    skip: int = Query(0),
    limit: int = Query(20),
    current_user: User = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db)
):
    await seed_hardware_if_empty(db)
    
    q_str = """
      SELECT hi.id, hi.category_id, hi.asset_code, hi.name, hi.brand, hi.model,
             hi.purchase_cost, hi.renting_price, hi.status,
             hi.pricing_unit, hi.description, hi.tax_category,
             hc.name as category_name,
             COALESCE(hs.available_quantity, 0) as inventory_count,
             COALESCE(hs.reserved_quantity, 0) as reserved_count,
             COUNT(*) OVER() as total_count
      FROM inventory.hardware_items hi
      LEFT JOIN inventory.hardware_categories hc ON hc.id = hi.category_id
      LEFT JOIN inventory.hardware_stock hs ON hs.hardware_id = hi.id
      WHERE (CAST(:category AS varchar) IS NULL OR hc.name ILIKE :cat_pct)
        AND (CAST(:status AS varchar) IS NULL OR hi.status = CAST(:status AS varchar))
        AND (CAST(:pricing_unit AS varchar) IS NULL OR hi.pricing_unit = CAST(:pricing_unit AS varchar))
        AND (CAST(:search AS varchar) IS NULL OR hi.name ILIKE :search_pct OR hi.asset_code ILIKE :search_pct)
      ORDER BY hi.asset_code ASC
      OFFSET :skip LIMIT :limit
    """
    params = {
        "category": category,
        "cat_pct": f"%{category}%" if category else None,
        "status": status,
        "pricing_unit": pricing_unit,
        "search": search,
        "search_pct": f"%{search}%" if search else None,
        "skip": skip,
        "limit": limit
    }
    
    rows_res = await db.execute(text(q_str), params)
    rows = rows_res.all()
    
    items = []
    total = 0
    for r in rows:
        total = r.total_count
        items.append({
            "id": str(r.id),
            "category_id": str(r.category_id),
            "item_code": r.asset_code,
            "asset_code": r.asset_code,
            "name": r.name,
            "brand": r.brand,
            "model": r.model,
            "cost_price": float(r.purchase_cost or 0),
            "purchase_cost": float(r.purchase_cost or 0),
            "selling_price": float(r.renting_price or 0),
            "renting_price": float(r.renting_price or 0),
            "status": r.status,
            "description": r.description or "",
            "tax_category": r.tax_category or "GST_18",
            "category_name": r.category_name,
            "pricing_unit": r.pricing_unit or "PER_EVENT",
            "inventory_count": r.inventory_count,
            "reserved_count": r.reserved_count,
            "is_active": r.status == "AVAILABLE"
        })
        
    sum_q = """
      SELECT
        COUNT(*) as total_items,
        COUNT(*) FILTER (WHERE status = 'AVAILABLE') as active_items,
        COUNT(DISTINCT category_id) as categories,
        AVG(purchase_cost) as avg_cost_price,
        AVG(renting_price) as avg_selling_price,
        SUM(purchase_cost * COALESCE(hs.quantity, 1)) as total_inventory_value
      FROM inventory.hardware_items hi
      LEFT JOIN inventory.hardware_stock hs ON hs.hardware_id = hi.id
    """
    sum_res = await db.execute(text(sum_q))
    s = sum_res.fetchone()
    
    summary = {
        "total_items": s.total_items or 0,
        "active_items": s.active_items or 0,
        "categories": s.categories or 0,
        "avg_cost_price": float(s.avg_cost_price or 0),
        "avg_selling_price": float(s.avg_selling_price or 0),
        "total_inventory_value": float(s.total_inventory_value or 0)
    }

    # Fetch categories that are actually assigned to at least one hardware item
    active_cats_res = await db.execute(
        select(HardwareCategory.name)
        .join(HardwareItem, HardwareItem.category_id == HardwareCategory.id)
        .distinct()
    )
    active_categories = [r[0] for r in active_cats_res.all() if r[0]]

    # Fetch pricing units that are actually assigned to at least one hardware item
    active_units_res = await db.execute(
        select(HardwareItem.pricing_unit)
        .distinct()
        .where(HardwareItem.pricing_unit.isnot(None))
    )
    active_pricing_units = [r[0] for r in active_units_res.all() if r[0]]
    
    # Calculate next hardware code sequentially
    next_item_code = "HW-1001"
    code_res = await db.execute(text("SELECT asset_code FROM inventory.hardware_items WHERE asset_code LIKE 'HW-%';"))
    all_codes = [r[0] for r in code_res.fetchall()]
    max_num = 1000
    for code in all_codes:
        if code.startswith("HW-"):
            try:
                num = int(code.split("-")[1])
                if num > max_num:
                    max_num = num
            except (ValueError, IndexError):
                pass
    next_item_code = f"HW-{max_num + 1}"
    
    return {
        "items": items,
        "total": total,
        "summary": summary,
        "next_item_code": next_item_code,
        "active_categories": sorted(active_categories),
        "active_pricing_units": sorted(active_pricing_units)
    }

@router.get("/superadmin/catalog/hardware/categories")
async def superadmin_get_hardware_categories(
    current_user: User = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db)
):
    await seed_hardware_if_empty(db)
    stmt = select(HardwareCategory).order_by(HardwareCategory.name)
    res = await db.execute(stmt)
    return [{"id": str(c.id), "name": c.name, "description": c.description} for c in res.scalars().all()]

@router.post("/superadmin/catalog/hardware")
async def superadmin_create_hardware(
    body: SuperAdminHardwareCreate,
    current_user: User = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db)
):
    code_str = body.item_code
    if not code_str or not code_str.strip():
        code_str = await get_next_available_hardware_code(db)
    else:
        existing_stmt = select(HardwareItem).where(HardwareItem.asset_code == code_str)
        existing = (await db.execute(existing_stmt)).scalar_one_or_none()
        if existing:
            code_str = await get_next_available_hardware_code(db)

    item = HardwareItem(
        id=uuid.uuid4(),
        category_id=body.category_id,
        asset_code=code_str,
        name=body.name,
        brand=body.brand or "Standard",
        model=body.model or "Generic v1",
        purchase_cost=body.cost_price,
        renting_price=body.selling_price,
        status="AVAILABLE",
        pricing_unit=body.pricing_unit or "PER_EVENT",
        description=body.description,
        tax_category=body.tax_category or "GST_18"
    )
    db.add(item)
    await db.flush()
    
    stock = HardwareStock(
        id=uuid.uuid4(),
        hardware_id=item.id,
        quantity=body.inventory_count or 0,
        reserved_quantity=0,
        available_quantity=body.inventory_count or 0
    )
    db.add(stock)
    await db.commit()
    
    return {"status": "success", "id": str(item.id)}

@router.patch("/superadmin/catalog/hardware/{item_id}")
async def superadmin_update_hardware(
    item_id: uuid.UUID,
    body: SuperAdminHardwareUpdate,
    current_user: User = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db)
):
    item_stmt = select(HardwareItem).where(HardwareItem.id == item_id)
    item = (await db.execute(item_stmt)).scalar_one_or_none()
    if not item:
        raise HTTPException(status_code=404, detail="Hardware item not found")
        
    if body.name is not None:
        item.name = body.name
    if body.item_code is not None:
        item.asset_code = body.item_code
    if body.category_id is not None:
        item.category_id = body.category_id
    if body.cost_price is not None:
        item.purchase_cost = body.cost_price
    if body.selling_price is not None:
        item.renting_price = body.selling_price
    if body.brand is not None:
        item.brand = body.brand
    if body.model is not None:
        item.model = body.model
    if body.pricing_unit is not None:
        item.pricing_unit = body.pricing_unit
    if body.description is not None:
        item.description = body.description
    if body.tax_category is not None:
        item.tax_category = body.tax_category
    if body.status is not None:
        item.status = body.status
        
    if body.inventory_count is not None:
        stock_stmt = select(HardwareStock).where(HardwareStock.hardware_id == item_id)
        stock = (await db.execute(stock_stmt)).scalar_one_or_none()
        
        qty = body.inventory_count
        if stock:
            stock.quantity = qty
            stock.available_quantity = qty - stock.reserved_quantity
        else:
            stock = HardwareStock(
                id=uuid.uuid4(),
                hardware_id=item_id,
                quantity=qty,
                reserved_quantity=0,
                available_quantity=qty
            )
            db.add(stock)
            
    await db.commit()
    return {"status": "success", "id": str(item_id)}

@router.post("/superadmin/catalog/hardware/import")
async def superadmin_import_hardware(
    file: UploadFile = File(...),
    current_user: User = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db)
):
    import io
    import re
    import openpyxl
    
    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {str(e)}")
        
    sheet = wb.active
    if sheet.max_row < 2:
        return {"status": "success", "count": 0}
        
    # Read headers
    header_cells = [cell.value for cell in sheet[1]]
    header_map = {}
    for idx, cell in enumerate(header_cells):
        if not cell:
            continue
        val = str(cell).strip().lower().replace("_", " ").replace("-", " ")
        header_map[val] = idx
        
    def get_val(row_vals, aliases, default=None):
        for alias in aliases:
            a_clean = alias.lower().replace("_", " ").replace("-", " ")
            if a_clean in header_map:
                res = row_vals[header_map[a_clean]]
                if res is not None:
                    return res
        return default

    # Load categories map
    cats_res = await db.execute(select(HardwareCategory))
    cats_list = cats_res.scalars().all()
    cat_map = {c.name.strip().lower(): c.id for c in cats_list}
    
    allocated_codes = set()
    imported_count = 0
    # Process rows
    for r_idx in range(2, sheet.max_row + 1):
        row_vals = [cell.value for cell in sheet[r_idx]]
        if not any(row_vals):
            continue # skip empty rows
            
        name = get_val(row_vals, ["hardware name", "name", "item name"])
        if not name or not str(name).strip():
            continue
            
        name_str = str(name).strip()
        item_code = get_val(row_vals, ["hardware code", "sku", "item code", "asset code"])
        if not item_code or not str(item_code).strip():
            code_str = await get_next_available_hardware_code(db, allocated_codes)
        else:
            item_code_str = str(item_code).strip()
            existing_stmt = select(HardwareItem).where(HardwareItem.asset_code == item_code_str)
            existing = (await db.execute(existing_stmt)).scalar_one_or_none()
            if existing or item_code_str in allocated_codes:
                code_str = await get_next_available_hardware_code(db, allocated_codes)
            else:
                code_str = item_code_str
                
        allocated_codes.add(code_str)
            
        # Resolve category
        cat_name = get_val(row_vals, ["category"], "Accessories")
        cat_name_clean = str(cat_name).strip().lower()
        if cat_name_clean in cat_map:
            cat_id = cat_map[cat_name_clean]
        else:
            # Create category dynamically
            cat_id = uuid.uuid4()
            new_cat = HardwareCategory(id=cat_id, name=str(cat_name).strip(), description="Created during bulk import")
            db.add(new_cat)
            await db.flush()
            cat_map[cat_name_clean] = cat_id
            
        pricing_unit = str(get_val(row_vals, ["pricing unit", "unit"], "PER_EVENT")).strip()
        if pricing_unit not in ["PER_DAY", "PER_EVENT", "PER_DEVICE", "PER_ROOM", "PER_COUNTER"]:
            pricing_unit = "PER_EVENT"
            
        cost_price = 0.0
        try:
            cost_price = float(get_val(row_vals, ["cost price", "cost"], 0.0))
        except ValueError:
            pass
            
        renting_price = 0.0
        try:
            renting_price = float(get_val(row_vals, ["renting price", "rent price", "selling price"], 0.0))
        except ValueError:
            pass
            
        brand = str(get_val(row_vals, ["brand"], "Standard")).strip()
        model = str(get_val(row_vals, ["model"], "Generic v1")).strip()
        
        inventory_count = 0
        try:
            inventory_count = int(get_val(row_vals, ["inventory count", "available quantity", "quantity", "inventory"], 0))
        except ValueError:
            pass
            
        description = get_val(row_vals, ["description"])
        desc_str = str(description).strip() if description else None
        
        tax_category = str(get_val(row_vals, ["tax category", "gst", "tax"], "GST_18")).strip()
        if tax_category not in ["GST_18", "GST_28", "GST_12", "GST_5", "GST_0"]:
            tax_category = "GST_18"
            
        # Create item
        item = HardwareItem(
            id=uuid.uuid4(),
            category_id=cat_id,
            asset_code=code_str,
            name=name_str,
            brand=brand,
            model=model,
            purchase_cost=cost_price,
            renting_price=renting_price,
            status="AVAILABLE",
            pricing_unit=pricing_unit,
            description=desc_str,
            tax_category=tax_category
        )
        db.add(item)
        await db.flush()
        
        # Create stock
        stock = HardwareStock(
            id=uuid.uuid4(),
            hardware_id=item.id,
            quantity=inventory_count,
            reserved_quantity=0,
            available_quantity=inventory_count
        )
        db.add(stock)
        imported_count += 1
        
    await db.commit()
    return {"status": "success", "count": imported_count}
