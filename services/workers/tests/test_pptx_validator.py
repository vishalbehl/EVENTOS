# workers/tests/test_pptx_validator.py
"""Tests for the PPTX validator lib."""

from __future__ import annotations

import io
import zipfile
import pytest


def _make_minimal_pptx() -> bytes:
    """Create a minimal, valid PPTX in memory using python-pptx."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        pytest.skip("python-pptx not installed")

    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    if title:
        title.text = "Test Slide"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_zip_part(data: bytes, name: str, content: bytes) -> bytes:
    source = io.BytesIO(data)
    out = io.BytesIO()
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(name, content)
    return out.getvalue()


class TestValidatePptx:
    def test_valid_pptx_returns_is_valid_true(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        assert result.is_valid is True
        assert result.slide_count >= 1
        assert result.errors == []

    def test_corrupt_pptx_returns_is_valid_false(self):
        from workers.lib.pptx_validator import validate_pptx
        result = validate_pptx(b"this is not a valid pptx file")
        assert result.is_valid is False
        assert len(result.errors) > 0

    def test_empty_bytes_returns_invalid(self):
        from workers.lib.pptx_validator import validate_pptx
        result = validate_pptx(b"")
        assert result.is_valid is False

    def test_slide_count_is_positive(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        assert result.slide_count >= 1

    def test_aspect_ratio_detected(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        # Default python-pptx slide is 16:9 or 10 x 7.5 (4:3)
        assert result.aspect_ratio is not None

    def test_no_external_links_on_clean_pptx(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        assert result.has_external_links is False
        assert result.external_link_urls == []

    def test_result_has_warnings_list(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        assert isinstance(result.warnings, list)

    def test_missing_fonts_is_list(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _make_minimal_pptx()
        result = validate_pptx(data)
        assert isinstance(result.missing_fonts, list)

    def test_macro_part_marks_file_invalid(self):
        from workers.lib.pptx_validator import validate_pptx
        data = _add_zip_part(_make_minimal_pptx(), "ppt/vbaProject.bin", b"macro")
        result = validate_pptx(data)
        assert result.has_macros is True
        assert result.is_valid is False

    def test_external_absolute_media_link_is_blocked(self):
        from workers.lib.pptx_validator import validate_pptx
        rels = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rRisk" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/video" Target="C:\\Users\\speaker\\Videos\\demo.mp4" TargetMode="External"/>
</Relationships>"""
        data = _add_zip_part(_make_minimal_pptx(), "ppt/slides/_rels/slide1.xml.rels", rels)
        result = validate_pptx(data)
        assert result.linked_assets_detected is True
        assert result.absolute_path_links_detected == 1
        assert result.is_valid is False


class TestPptxValidationResult:
    def test_dataclass_defaults(self):
        from workers.lib.pptx_validator import PptxValidationResult
        r = PptxValidationResult()
        assert r.is_valid is True
        assert r.slide_count == 0
        assert r.errors == []
        assert r.warnings == []
        assert r.missing_fonts == []
        assert r.external_link_urls == []
